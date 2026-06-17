from pathlib import Path

from docx import Document
from docx.shared import Pt
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt as PptPt


BASE = Path(r"E:/diploma/outputs/019ed18e-8ec4-7a51-b7bc-1a735bc7d652/presentations/vkr-defense")
TPL = BASE / "template.pptx"
AN = BASE / "analysis"
OUT = BASE / "output"
OUT.mkdir(parents=True, exist_ok=True)

PPTX_OUT = OUT / "Zozin_Puls_VKR_presentation.pptx"
DOCX_OUT = OUT / "Zozin_Puls_text_zashchity.docx"

BLUE = RGBColor(0x2E, 0x5A, 0xAC)
DARK_BLUE = RGBColor(0x32, 0x3C, 0x8D)
RED = RGBColor(0xC0, 0x39, 0x2B)
GREEN = RGBColor(0x1F, 0x9D, 0x72)
GRAY = RGBColor(0x5B, 0x64, 0x76)
LIGHT = RGBColor(0xEA, 0xEC, 0xF6)
BLACK = RGBColor(0x1F, 0x26, 0x37)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)


def I(v):
    return Inches(v)


def remove_shape(shape):
    shape._element.getparent().remove(shape._element)


def clear_body(slide):
    for shape in list(slide.shapes):
        top = shape.top
        if I(1.0) < top < I(6.65):
            remove_shape(shape)


def set_text(shape, text, size=18, color=BLACK, bold=False, align=PP_ALIGN.LEFT):
    shape.text_frame.clear()
    shape.text_frame.word_wrap = True
    shape.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = shape.text_frame.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = PptPt(size)
    run.font.bold = bold
    run.font.name = "Arial"
    run.font.color.rgb = color


def add_text(slide, x, y, w, h, text, size=18, color=BLACK, bold=False, align=PP_ALIGN.LEFT):
    shape = slide.shapes.add_textbox(I(x), I(y), I(w), I(h))
    set_text(shape, text, size, color, bold, align)
    return shape


def card(slide, x, y, w, h, fill=WHITE, line=LIGHT, radius=True):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, I(x), I(y), I(w), I(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    shp.line.color.rgb = line
    shp.line.width = PptPt(1)
    return shp


def pill(slide, x, y, w, h, text, fill=BLUE, color=WHITE, size=18):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, I(x), I(y), I(w), I(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    shp.line.color.rgb = fill
    set_text(shp, text, size=size, color=color, bold=True, align=PP_ALIGN.CENTER)
    return shp


def set_title(slide, title, num=None):
    for shape in slide.shapes:
        if shape.top < I(0.8) and hasattr(shape, "text") and shape.text.strip():
            set_text(shape, title, size=32, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
            break
    if num is not None:
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.left > I(11.5) and shape.top > I(6.6):
                set_text(shape, str(num), size=14, color=GRAY, align=PP_ALIGN.CENTER)


def bullet_lines(shape, lines, size=16, color=BLACK):
    shape.text_frame.clear()
    shape.text_frame.word_wrap = True
    for i, line in enumerate(lines):
        p = shape.text_frame.paragraphs[0] if i == 0 else shape.text_frame.add_paragraph()
        p.text = line
        p.level = 0
        p.font.size = PptPt(size)
        p.font.name = "Arial"
        p.font.color.rgb = color
        p.space_after = PptPt(3)


def add_bullets(slide, x, y, w, h, lines, size=16):
    shape = slide.shapes.add_textbox(I(x), I(y), I(w), I(h))
    bullet_lines(shape, lines, size=size)
    return shape


def bar_chart(slide, x, y, w, h, labels, values, max_value=None, color=BLUE, suffix=""):
    max_value = max_value or max(values)
    axis_y = y + h - 0.35
    add_text(slide, x, y - 0.08, w, 0.25, "Расчетная оценка времени основных сценариев", size=13, color=GRAY, bold=True)
    for i, (label, value) in enumerate(zip(labels, values)):
        bx = x + i * (w / len(labels)) + 0.18
        bw = (w / len(labels)) - 0.36
        bh = (h - 0.85) * value / max_value
        card(slide, bx, axis_y - bh, bw, bh, fill=color, line=color, radius=False)
        add_text(slide, bx, axis_y - bh - 0.34, bw, 0.25, f"{value}{suffix}", size=15, color=color, bold=True, align=PP_ALIGN.CENTER)
        add_text(slide, bx - 0.05, axis_y + 0.05, bw + 0.1, 0.45, label, size=10.5, color=GRAY, align=PP_ALIGN.CENTER)
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, I(x), I(axis_y), I(w), I(0.01))
    line.fill.solid()
    line.fill.fore_color.rgb = LIGHT
    line.line.color.rgb = LIGHT


def mini_bar(slide, x, y, w, label, value, max_value, color=BLUE):
    add_text(slide, x, y, 2.3, 0.28, label, size=12.5, color=BLACK)
    card(slide, x + 2.35, y + 0.04, w - 2.7, 0.16, fill=LIGHT, line=LIGHT, radius=True)
    card(slide, x + 2.35, y + 0.04, (w - 2.7) * value / max_value, 0.16, fill=color, line=color, radius=True)
    add_text(slide, x + w - 0.35, y - 0.03, 0.45, 0.25, str(value), size=12, color=color, bold=True, align=PP_ALIGN.RIGHT)


def add_picture_fit(slide, path, x, y, w, h):
    pic = slide.shapes.add_picture(str(path), I(x), I(y), width=I(w))
    if pic.height > I(h):
        pic.height = I(h)
    if pic.width > I(w):
        pic.width = I(w)
    pic.left = I(x) + int((I(w) - pic.width) / 2)
    pic.top = I(y) + int((I(h) - pic.height) / 2)
    return pic


prs = Presentation(str(TPL))

# Slide 1
s = prs.slides[0]
set_text(s.shapes[0], "Министерство науки и высшего образования Российской Федерации\nФГБОУ ВО «Магнитогорский государственный технический университет им. Г.И. Носова»\nИнститут энергетики и автоматизированных систем · Кафедра ВТиП", 14, GRAY, align=PP_ALIGN.CENTER)
set_text(s.shapes[1], "ПРОЕКТИРОВАНИЕ И РАЗРАБОТКА ИНФОРМАЦИОННОЙ СИСТЕМЫ УПРАВЛЕНИЯ ВОЛОНТЕРАМИ С ЭЛЕМЕНТАМИ ГЕЙМИФИКАЦИИ И АНАЛИТИКИ", 25, DARK_BLUE, True, PP_ALIGN.CENTER)
set_text(s.shapes[2], "Выполнил: обучающийся Зозин П.А.\nНаправление 09.03.01 — Информатика и вычислительная техника\nРуководитель: доцент кафедры ВТиП, к.п.н. Гладышева М.М.", 14, BLACK, align=PP_ALIGN.RIGHT)
set_text(s.shapes[3], "Магнитогорск, 2026", 14, GRAY, align=PP_ALIGN.RIGHT)

# Slide 2
s = prs.slides[1]
clear_body(s)
set_title(s, "Цель, объект, предмет и задачи работы", 2)
card(s, 0.7, 1.25, 11.9, 1.25, fill=DARK_BLUE, line=DARK_BLUE)
add_text(s, 0.95, 1.38, 1.45, 0.35, "ЦЕЛЬ", 19, WHITE, True)
add_text(s, 2.1, 1.28, 10.0, 0.85, "Повысить эффективность управления волонтерской деятельностью за счет разработки web-ориентированной информационной системы «Пульс» с учетом заявок, задач, часов, сертификатов, геймификации и аналитики.", 18, WHITE)
card(s, 0.7, 2.75, 5.75, 0.95)
add_text(s, 0.95, 2.92, 1.25, 0.35, "ОБЪЕКТ", 15, BLUE, True)
add_text(s, 2.05, 2.83, 4.05, 0.55, "процесс управления волонтерской деятельностью в организации", 15, BLACK)
card(s, 6.75, 2.75, 5.85, 0.95)
add_text(s, 7.0, 2.92, 1.45, 0.35, "ПРЕДМЕТ", 15, BLUE, True)
add_text(s, 8.35, 2.83, 3.95, 0.55, "информационная система «Пульс» и связанные пользовательские сценарии", 15, BLACK)
tasks = [
    "Проанализировать предметную область и аналоги",
    "Спроектировать архитектуру, БД и REST API",
    "Реализовать backend, frontend и роли доступа",
    "Связать заявки, часы, достижения и сертификаты",
    "Проверить сборку и ключевые сценарии",
    "Оценить практический эффект и план пилота",
]
for i, t in enumerate(tasks):
    x = 0.7 + (i % 2) * 6.05
    y = 4.0 + (i // 2) * 0.78
    card(s, x, y, 5.75, 0.58)
    pill(s, x + 0.15, y + 0.11, 0.38, 0.34, str(i + 1), BLUE, WHITE, 13)
    add_text(s, x + 0.65, y + 0.08, 4.85, 0.38, t, 13.8, BLACK)

# Slide 3
s = prs.slides[2]
clear_body(s)
set_title(s, "Анализ аналогов: нужен локальный полный контур", 3)
analogs = [
    ("D", "Добро.рф", "масштабная публичная экосистема, но не локальное пространство организации"),
    ("I/V", "Idealist / VolunteerMatch", "сильный поиск возможностей, но слабый учет внутреннего вклада"),
    ("T", "Timecounts", "операционная работа событий, но нужна адаптация под аналитику и сертификаты"),
    ("G", "Golden", "коммерческая SaaS-платформа, меньше контроля над моделью данных"),
]
for i, (abbr, name, note) in enumerate(analogs):
    x = 0.7 + i * 3.07
    card(s, x, 1.35, 2.7, 3.0)
    pill(s, x + 0.78, 1.65, 1.12, 0.72, abbr, [BLUE, DARK_BLUE, RED, BLACK][i], WHITE, 21)
    add_text(s, x + 0.2, 2.55, 2.3, 0.42, name, 15.5, BLACK, True, PP_ALIGN.CENTER)
    add_text(s, x + 0.28, 3.08, 2.15, 0.9, note, 11.6, GRAY, align=PP_ALIGN.CENTER)
card(s, 0.7, 4.63, 11.9, 1.05, fill=DARK_BLUE, line=DARK_BLUE)
add_text(s, 0.95, 4.75, 10.75, 0.55, "Вывод: существующие решения закрывают отдельные части процесса, но не связывают локальные роли, заявки, подтвержденные часы, достижения, сертификаты и аналитику в одном контуре данных.", 16.5, WHITE)
add_text(s, 0.8, 5.92, 2.5, 0.25, "Покрытие полного цикла", 12, GRAY, True)
for j, (label, val) in enumerate([("каталог", 2), ("event", 3), ("SaaS", 4), ("Пульс", 6)]):
    bx = 3.3 + j * 2.15
    add_text(s, bx, 5.88, 0.85, 0.25, label, size=11, color=GRAY)
    card(s, bx + 0.78, 5.96, 0.9, 0.13, fill=LIGHT, line=LIGHT, radius=True)
    card(s, bx + 0.78, 5.96, 0.9 * val / 6, 0.13, fill=GREEN if label == "Пульс" else BLUE, line=GREEN if label == "Пульс" else BLUE, radius=True)
    add_text(s, bx + 1.72, 5.86, 0.22, 0.25, str(val), size=11, color=GREEN if label == "Пульс" else BLUE, bold=True)

# Slide 4
s = prs.slides[3]
clear_body(s)
set_title(s, "Технологический стек", 4)
cols = [
    ("Frontend", BLUE, ["Vue.js 3", "TypeScript", "Vite", "страницы ролей", "адаптивный интерфейс"]),
    ("Backend", DARK_BLUE, ["Go + Gin", "REST API", "JWT / refresh-сессии", "сервисы и репозитории", "audit log"]),
    ("Данные / DevOps", RED, ["PostgreSQL", "миграции", "Docker Compose", "файловые загрузки", "сертификаты и аналитика"]),
]
for i, (title, color, items) in enumerate(cols):
    x = 0.7 + i * 4.03
    card(s, x, 1.42, 3.55, 4.95)
    card(s, x, 1.42, 3.55, 0.68, fill=color, line=color, radius=False)
    add_text(s, x, 1.52, 3.55, 0.42, title, 18, WHITE, True, PP_ALIGN.CENTER)
    for j, item in enumerate(items):
        pill(s, x + 0.32, 2.38 + j * 0.68, 0.28, 0.28, "", color, WHITE, 8)
        add_text(s, x + 0.78, 2.25 + j * 0.68, 2.3, 0.42, item, 16, BLACK)
add_text(s, 0.8, 6.25, 11.2, 0.32, "Стек выбран под web-приложение с разграничением ролей, воспроизводимым развертыванием и реляционной целостностью данных.", 14, GRAY, align=PP_ALIGN.CENTER)

# Slide 5
s = prs.slides[4]
clear_body(s)
set_title(s, "Архитектура «Пульс» связывает роли, API и данные", 5)
layers = [
    ("Уровень представления", "Vue.js SPA", "личный кабинет волонтера, панели координатора, администратора организации и суперадминистратора"),
    ("Уровень бизнес-логики", "Go / Gin REST API", "авторизация, организации, мероприятия, заявки, задачи, часы, достижения, сертификаты, аналитика"),
    ("Уровень данных", "PostgreSQL", "пользователи, роли, организации, события, time entries, audit log и агрегаты отчетности"),
]
for i, (a, b, c) in enumerate(layers):
    y = 1.35 + i * 1.35
    card(s, 0.7, y, 11.9, 1.05)
    card(s, 0.7, y, 3.35, 1.05, fill=[BLUE, DARK_BLUE, RED][i], line=[BLUE, DARK_BLUE, RED][i], radius=False)
    add_text(s, 1.0, y + 0.17, 2.55, 0.28, a, 16.5, WHITE, True)
    add_text(s, 1.0, y + 0.57, 2.55, 0.25, b, 13, WHITE)
    add_text(s, 4.35, y + 0.17, 7.55, 0.55, c, 15, BLACK)
    if i < 2:
        pill(s, 6.25, y + 1.12, 0.62, 0.36, "↓", GREEN, WHITE, 15)
card(s, 0.7, 5.75, 11.9, 0.62, fill=DARK_BLUE, line=DARK_BLUE)
add_text(s, 0.95, 5.83, 10.9, 0.35, "Ключевой принцип: заявка, посещаемость, подтвержденные часы, достижение, сертификат и KPI не разрознены, а проходят через единую модель данных.", 14.5, WHITE)

# Slide 6
s = prs.slides[5]
clear_body(s)
set_title(s, "Модель данных покрывает обязательный контур управления", 6)
domains = [
    ("Пользователи и роли", "users, system_roles, memberships"),
    ("Организации и события", "organizations, events, applications"),
    ("Работа и вклад", "tasks, attendance, time_entries"),
    ("Мотивация и отчетность", "achievements, certificates, analytics, audit_log"),
]
for i, (name, desc) in enumerate(domains):
    y = 1.28 + i * 1.0
    card(s, 0.7, y, 5.95, 0.75)
    add_text(s, 0.95, y + 0.12, 2.25, 0.28, name, 14.5, BLUE, True)
    add_text(s, 3.15, y + 0.08, 3.1, 0.36, desc, 11.8, GRAY)
add_text(s, 0.85, 5.35, 5.55, 0.3, "Физическая модель: 18+ предметных таблиц, внешние ключи, статусы и журнал аудита.", 13.5, BLACK, True)
add_picture_fit(s, AN / "image34.png", 7.05, 1.25, 5.05, 4.25)
card(s, 7.05, 5.65, 5.05, 0.58, fill=DARK_BLUE, line=DARK_BLUE)
add_text(s, 7.25, 5.73, 4.65, 0.32, "ER-диаграмма ядра: данные волонтера связаны с организацией, заявками, задачами и подтвержденным вкладом.", 12.2, WHITE)

# Slide 7
s = prs.slides[6]
clear_body(s)
set_title(s, "Ключевой сценарий проходит от заявки до управленческого отчета", 7)
steps = [
    ("1", "Волонтер", "подает заявку на мероприятие"),
    ("2", "Координатор", "рассматривает заявку и фиксирует участие"),
    ("3", "Система", "связывает мероприятие, задачу и запись времени"),
    ("4", "Координатор", "подтверждает часы по основанию"),
    ("5", "Пульс", "начисляет достижения и формирует сертификат"),
    ("6", "Руководитель", "видит KPI и выгружает отчет"),
]
for i, (num, role, action) in enumerate(steps):
    x = 0.75 + (i % 2) * 6.0
    y = 1.35 + (i // 2) * 1.42
    card(s, x, y, 5.6, 1.02)
    pill(s, x + 0.22, y + 0.25, 0.52, 0.52, num, [BLUE, DARK_BLUE, RED, BLUE, GREEN, DARK_BLUE][i], WHITE, 17)
    add_text(s, x + 0.95, y + 0.18, 1.65, 0.25, role, 14.5, BLUE, True)
    add_text(s, x + 0.95, y + 0.47, 4.15, 0.3, action, 13.5, BLACK)
card(s, 0.75, 5.93, 11.55, 0.52, fill=DARK_BLUE, line=DARK_BLUE)
add_text(s, 1.0, 6.0, 10.95, 0.28, "Главный результат сценария: управленческие показатели строятся на подтвержденных операционных данных, а не на ручном сведении таблиц.", 13.5, WHITE)

# Slide 8
s = prs.slides[7]
clear_body(s)
set_title(s, "Реализованный интерфейс закрывает основные роли", 8)
imgs = [
    ("image29.jpg", "вход"),
    ("image28.jpg", "панель"),
    ("image25.jpg", "мероприятия"),
    ("image14.jpg", "задачи"),
    ("image19.jpg", "часы"),
    ("image13.jpg", "сертификаты"),
]
for i, (img, label) in enumerate(imgs):
    x = 0.75 + (i % 3) * 4.05
    y = 1.23 + (i // 3) * 2.45
    card(s, x, y, 3.65, 2.08)
    add_picture_fit(s, AN / img, x + 0.12, y + 0.12, 3.41, 1.62)
    add_text(s, x + 0.18, y + 1.78, 3.25, 0.22, label, 11.5, BLUE, True, PP_ALIGN.CENTER)
add_text(s, 0.85, 6.22, 11.4, 0.25, "Экранные формы показывают, что результат работы не ограничен схемами: реализованы рабочие разделы для заявок, задач, часов, сертификатов и администрирования.", 12.5, GRAY, align=PP_ALIGN.CENTER)

# Slide 9
s = prs.slides[8]
clear_body(s)
set_title(s, "Проверка работоспособности: что подтверждено", 9)
kpis = [
    ("npm build", "успешно", "1782 модуля Vite"),
    ("go test ./...", "сборка OK", "автотестов нет"),
    ("Профиль", "проверен", "миграция 020 + API"),
    ("Аватар", "проверен", "upload PNG"),
]
for i, (a, b, c) in enumerate(kpis):
    x = 0.75 + i * 3.02
    card(s, x, 1.35, 2.65, 1.16)
    add_text(s, x + 0.15, 1.48, 2.35, 0.28, a, 12.5, GRAY, True, PP_ALIGN.CENTER)
    add_text(s, x + 0.15, 1.76, 2.35, 0.35, b, 18, GREEN if i in [0, 2, 3] else BLUE, True, PP_ALIGN.CENTER)
    add_text(s, x + 0.15, 2.13, 2.35, 0.22, c, 10.8, GRAY, align=PP_ALIGN.CENTER)
card(s, 0.75, 2.95, 5.75, 2.75)
add_text(s, 1.0, 3.12, 5.2, 0.3, "Матрица сценариев", 15, BLUE, True)
checks = [("подтверждено запуском", 4, GREEN), ("подтверждено реализацией", 2, BLUE), ("требует опытной проверки", 2, RED)]
for i, (label, val, col) in enumerate(checks):
    mini_bar(s, 1.05, 3.65 + i * 0.55, 4.9, label, val, 4, col)
add_text(s, 1.0, 5.45, 5.1, 0.35, "Формулировка результата остается корректной: подтверждены сборка, часть API и реализованный контур; полный приемочный вывод переносится в пилот.", 11.7, GRAY)
card(s, 6.85, 2.95, 5.45, 2.75)
add_text(s, 7.1, 3.12, 4.9, 0.3, "Ограничения закрываются планом пилота", 15, BLUE, True)
add_bullets(s, 7.1, 3.58, 4.9, 1.45, [
    "чистая PostgreSQL и миграции",
    "сценарий заявка -> часы -> аналитика -> сертификат",
    "unit/integration tests прав доступа",
    "измерение времени на контрольных сценариях",
], 12.5)

# Slide 10
s = prs.slides[9]
clear_body(s)
set_title(s, "Практический эффект: меньше ручного сведения и быстрее сценарии", 10)
bar_chart(s, 0.8, 1.45, 5.6, 3.0, ["заявка", "часы", "KPI"], [35, 45, 50], 50, BLUE, " с")
card(s, 6.85, 1.35, 5.35, 3.05)
add_text(s, 7.1, 1.55, 4.85, 0.28, "Что заменяет система", 15, BLUE, True)
rows = [
    ("Сбор заявок", "статусная заявка в системе"),
    ("Учет часов", "запись связана с задачей/событием"),
    ("Сертификат", "формируется по подтвержденным данным"),
    ("Отчетность", "KPI строятся из операционной БД"),
]
for i, (a, b) in enumerate(rows):
    y = 2.05 + i * 0.48
    add_text(s, 7.1, y, 1.55, 0.22, a, 11.5, BLACK, True)
    add_text(s, 8.75, y, 3.05, 0.25, b, 11.5, GRAY)
card(s, 0.8, 4.85, 11.4, 1.15, fill=DARK_BLUE, line=DARK_BLUE)
add_text(s, 1.05, 5.02, 10.8, 0.5, "Практическая значимость: «Пульс» можно использовать как основу автоматизации малых и средних волонтерских организаций, студенческих объединений и социальных проектов.", 17, WHITE, align=PP_ALIGN.CENTER)

# Slide 11
s = prs.slides[10]
set_text(s.shapes[0], "Спасибо за внимание!", 42, DARK_BLUE, True, PP_ALIGN.CENTER)
set_text(s.shapes[1], "Информационная система «Пульс»: заявки, задачи, часы, достижения, сертификаты и аналитика в едином контуре", 18, GRAY, align=PP_ALIGN.CENTER)

prs.save(PPTX_OUT)


script = [
    ("Слайд 1. Титульный", "Уважаемые члены государственной экзаменационной комиссии, вашему вниманию представляется выпускная квалификационная работа на тему «Проектирование и разработка информационной системы управления волонтерами с элементами геймификации и аналитики». Работа выполнена на примере системы «Пульс»."),
    ("Слайд 2. Цель и задачи", "Цель работы состоит в повышении эффективности управления волонтерской деятельностью за счет разработки web-ориентированной информационной системы. Объектом является процесс управления волонтерской деятельностью в организации, предметом - система «Пульс» и связанные пользовательские сценарии. Для достижения цели были решены задачи анализа предметной области, проектирования архитектуры и базы данных, реализации серверной и клиентской частей, а также проверки ключевых сценариев."),
    ("Слайд 3. Анализ аналогов", "На этапе анализа были рассмотрены Добро.рф, Idealist / VolunteerMatch, Timecounts и Golden. Эти решения полезны, но закрывают разные части процесса. Одни ориентированы на публичный поиск волонтерских возможностей, другие - на event-management или коммерческую SaaS-модель. Для дипломного проекта важен локальный контур организации, где заявка, участие, часы, достижения, сертификаты и аналитика связаны между собой."),
    ("Слайд 4. Технологический стек", "Система реализована как web-приложение. Клиентская часть построена на Vue.js 3, TypeScript и Vite. Серверная часть реализована на Go с REST API, авторизацией и разделением ролей. Для хранения данных используется PostgreSQL, а воспроизводимость развертывания поддерживается Docker-конфигурацией."),
    ("Слайд 5. Архитектура", "Архитектура включает три уровня: представление, бизнес-логику и данные. Пользователь работает через SPA-интерфейс. Серверный API обрабатывает авторизацию, организации, мероприятия, заявки, задачи, часы, сертификаты и аналитику. В базе данных сохраняются связанные сущности, поэтому управленческие показатели строятся не вручную, а на подтвержденных операционных данных."),
    ("Слайд 6. Модель данных", "Модель данных покрывает пользователей и роли, организации и события, работу и вклад волонтера, а также мотивацию и отчетность. В дипломе приведены ER-диаграммы и физическое описание таблиц. Важная особенность - целостность связей между заявками, посещаемостью, временем, достижениями и аудитом действий."),
    ("Слайд 7. Ключевой сценарий", "Ключевой пользовательский сценарий начинается с подачи заявки волонтером. Координатор рассматривает заявку и фиксирует участие. Затем система связывает мероприятие, задачу и запись времени, после чего подтвержденные часы становятся основанием для достижений, сертификата и аналитики. Так формируется единый путь от действия пользователя до управленческого отчета."),
    ("Слайд 8. Интерфейс", "В работе реализованы основные экранные формы: вход, главная панель, мероприятия, задачи, учет часов, сертификаты и административные разделы. Интерфейс на текущем этапе функциональный, что также отражено в отзыве: визуальное оформление можно развивать дальше, но основные пользовательские сценарии уже представлены в приложении."),
    ("Слайд 9. Проверка", "Проверка работоспособности включала сборку frontend и backend, проверку миграции и API расширенного профиля, а также загрузку аватара. При этом в работе честно зафиксированы ограничения: автоматизированные unit и integration tests отсутствуют, а часть KPI требует опытной эксплуатации. Поэтому предложен план закрытия рисков через запуск миграций на чистой базе, сценарный проход и добавление тестов прав доступа."),
    ("Слайд 10. Практический эффект", "Практический эффект состоит в сокращении ручных операций координатора. Вместо отдельных форм, таблиц и переписок система хранит заявку, часы, сертификат и показатели отчетности в единой базе. Расчетные оценки основных сценариев показывают, что типовые действия можно выполнить в пределах 35-50 секунд, а полный фактический замер должен быть выполнен на пилоте."),
    ("Слайд 11. Завершение", "В результате работы разработана информационная система «Пульс», которая объединяет управление волонтерами, мероприятиями, задачами, подтвержденными часами, достижениями, сертификатами и аналитикой. Система может использоваться как основа автоматизации малых и средних волонтерских организаций, студенческих объединений и социальных проектов. Спасибо за внимание, я готов ответить на вопросы."),
]

doc = Document()
styles = doc.styles
styles["Normal"].font.name = "Arial"
styles["Normal"].font.size = Pt(11)
doc.add_heading("Текст выступления к защите ВКР", level=1)
doc.add_paragraph("Зозин П.А. Тема: «Проектирование и разработка информационной системы управления волонтерами с элементами геймификации и аналитики».")
for title, text in script:
    doc.add_heading(title, level=2)
    doc.add_paragraph(text)
doc.add_heading("Короткая версия вступления", level=2)
doc.add_paragraph("Уважаемые члены комиссии, в работе разработана информационная система «Пульс» для управления волонтерской деятельностью. Система решает проблему разрозненного учета заявок, задач, часов и отчетности, объединяя их в одном web-приложении с ролями, геймификацией, сертификатами и аналитикой.")
doc.add_heading("Ключевые акценты для ответов", level=2)
for item in [
    "Новизна для проекта: не каталог мероприятий, а локальный контур организации с подтвержденным вкладом и отчетностью.",
    "Основной результат: реализованы backend, frontend, база данных и ключевые модули «Пульс».",
    "Честное ограничение: полная приемка и KPI требуют опытной эксплуатации; это закрывается предложенным планом пилота.",
    "Практическая значимость: система подходит как основа автоматизации для небольших организаций и студенческих объединений.",
]:
    doc.add_paragraph(item, style=None)
doc.save(DOCX_OUT)

print(PPTX_OUT)
print(DOCX_OUT)
