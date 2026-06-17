import json, zipfile, re
from pathlib import Path
from docx import Document

docx_path = Path(r'E:/diploma/outputs/019ed18e-8ec4-7a51-b7bc-1a735bc7d652/presentations/vkr-defense/source.docx')
out_dir = Path(r'E:/diploma/outputs/019ed18e-8ec4-7a51-b7bc-1a735bc7d652/presentations/vkr-defense/analysis')
out_dir.mkdir(parents=True, exist_ok=True)
doc = Document(str(docx_path))
items=[]
for i,p in enumerate(doc.paragraphs,1):
    text = p.text.strip()
    if text:
        items.append({'type':'paragraph','index':i,'style':p.style.name if p.style else '', 'text': text})
for ti,t in enumerate(doc.tables,1):
    rows=[]
    for row in t.rows:
        rows.append([cell.text.strip() for cell in row.cells])
    items.append({'type':'table','index':ti,'rows':rows})
headings=[]; captions=[]; formulas=[]
ru_prefixes = ('Рисунок','Таблица','Диаграмма','График')
for it in items:
    if it['type']=='paragraph':
        text=it['text']
        style=it.get('style','')
        if 'Heading' in style or 'Заголов' in style or re.match(r'^[0-9]+(\.[0-9]+)*\s+', text):
            headings.append(it)
        if text.startswith(ru_prefixes):
            captions.append(it)
        if any(s in text.lower() for s in ['=', 'формул', 'коэффициент', 'показател']):
            formulas.append(it)
media=[]
with zipfile.ZipFile(docx_path) as z:
    for n in z.namelist():
        if n.startswith('word/media/'):
            data=z.read(n)
            media.append({'name': n, 'size': len(data)})
            (out_dir / Path(n).name).write_bytes(data)
full_text=[]
for it in items:
    if it['type']=='paragraph':
        full_text.append(it['text'])
    else:
        full_text.append('\n'.join(['\t'.join(r) for r in it['rows']]))
(out_dir/'diploma_full_text.txt').write_text('\n\n'.join(full_text), encoding='utf-8')
(out_dir/'diploma_extracted.json').write_text(json.dumps({'items':items,'headings':headings,'captions':captions,'formulas':formulas[:200],'media':media}, ensure_ascii=False, indent=2), encoding='utf-8')
print(json.dumps({'paragraphs': len(doc.paragraphs), 'tables': len(doc.tables), 'nonempty_items': len(items), 'headings': len(headings), 'captions': len(captions), 'media': len(media), 'out': str(out_dir)}, ensure_ascii=False, indent=2))
