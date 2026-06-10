from __future__ import annotations

import io
import shutil
import zipfile
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont


FINAL_DIR = Path(r"E:\diploma\docs\diploma_document\zozin_final_diploma_files")
INPUT_DOCX = FINAL_DIR / "Zozin_v8.docx"
REFERENCE_DOCX = FINAL_DIR / "Zozin_v7.docx"
OUTPUT_DOCX = FINAL_DIR / "Zozin_v9.docx"
WORK_DIR = FINAL_DIR / "_codex_service" / "generated" / "V9_assets"
TEXT_STAGE = WORK_DIR / "Zozin_v9_text_stage.docx"

EMU_PER_INCH = 914400


def font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont:
    candidates = [
        r"C:\Windows\Fonts\arialbd.ttf" if bold else r"C:\Windows\Fonts\arial.ttf",
        r"C:\Windows\Fonts\calibrib.ttf" if bold else r"C:\Windows\Fonts\calibri.ttf",
    ]
    for path in candidates:
        if Path(path).exists():
            return ImageFont.truetype(path, size)
    return ImageFont.load_default()


def draw_centered_text(draw: ImageDraw.ImageDraw, box: tuple[int, int, int, int], text: str, fnt: ImageFont.FreeTypeFont) -> None:
    words = text.split()
    lines: list[str] = []
    current = ""
    max_width = box[2] - box[0] - 28
    for word in words:
        candidate = f"{current} {word}".strip()
        if draw.textbbox((0, 0), candidate, font=fnt)[2] <= max_width:
            current = candidate
        else:
            if current:
                lines.append(current)
            current = word
    if current:
        lines.append(current)

    line_heights = [draw.textbbox((0, 0), line, font=fnt)[3] for line in lines]
    total_height = sum(line_heights) + max(0, len(lines) - 1) * 6
    y = box[1] + ((box[3] - box[1]) - total_height) // 2
    for line, h in zip(lines, line_heights):
        w = draw.textbbox((0, 0), line, font=fnt)[2]
        draw.text((box[0] + ((box[2] - box[0]) - w) // 2, y), line, font=fnt, fill=(0, 0, 0))
        y += h + 6


def rounded_box(draw: ImageDraw.ImageDraw, box: tuple[int, int, int, int], text: str, fill: tuple[int, int, int], fnt: ImageFont.FreeTypeFont) -> None:
    draw.rounded_rectangle(box, radius=14, fill=fill, outline=(20, 20, 20), width=3)
    draw_centered_text(draw, box, text, fnt)


def arrow(draw: ImageDraw.ImageDraw, start: tuple[int, int], end: tuple[int, int]) -> None:
    draw.line([start, end], fill=(20, 20, 20), width=3)
    x1, y1 = start
    x2, y2 = end
    if abs(y2 - y1) >= abs(x2 - x1):
        direction = 1 if y2 >= y1 else -1
        pts = [(x2, y2), (x2 - 10, y2 - 18 * direction), (x2 + 10, y2 - 18 * direction)]
    else:
        direction = 1 if x2 >= x1 else -1
        pts = [(x2, y2), (x2 - 18 * direction, y2 - 10), (x2 - 18 * direction, y2 + 10)]
    draw.polygon(pts, fill=(20, 20, 20))


def create_general_screen_map(path: Path) -> None:
    img = Image.new("RGB", (1600, 1180), "white")
    draw = ImageDraw.Draw(img)
    node_font = font(25)
    small_font = font(22)

    rounded_box(draw, (560, 40, 1040, 120), "Вход / авторизация", (242, 245, 248), node_font)

    role_boxes = {
        "Волонтер": (70, 285, 360, 375),
        "Координатор": (430, 285, 730, 375),
        "Администратор организации": (790, 285, 1120, 375),
        "Суперадминистратор": (1180, 285, 1530, 375),
    }
    role_centers = [((box[0] + box[2]) // 2, box[1]) for box in role_boxes.values()]
    bus_y = 210
    draw.line([(800, 120), (800, bus_y)], fill=(20, 20, 20), width=3)
    draw.line([(role_centers[0][0], bus_y), (role_centers[-1][0], bus_y)], fill=(20, 20, 20), width=3)
    for text, box in role_boxes.items():
        rounded_box(draw, box, text, (242, 245, 248), small_font)
        arrow(draw, ((box[0] + box[2]) // 2, bus_y), ((box[0] + box[2]) // 2, box[1]))

    columns = {
        "Волонтер": [
            "Профиль и уведомления",
            "Мероприятия и заявки",
            "Мои задачи",
            "Учет часов",
            "Достижения и сертификаты",
        ],
        "Координатор": [
            "Мероприятия",
            "Заявки участников",
            "Задачи и назначения",
            "Посещаемость",
            "Подтверждение часов",
        ],
        "Администратор организации": [
            "Пользователи и роли",
            "Настройки организации",
            "Поля профиля",
            "Контент и база знаний",
            "Аналитика и выгрузки",
        ],
        "Суперадминистратор": [
            "Системные роли",
            "Организации",
            "Справочники",
            "Журнал аудита",
            "Системные настройки",
        ],
    }
    x_positions = [70, 430, 790, 1180]
    for col_idx, (role, items) in enumerate(columns.items()):
        x = x_positions[col_idx]
        prev_bottom = 375
        center_x = x + 145 if col_idx < 2 else x + 165
        width = 290 if col_idx < 2 else 350
        for row_idx, item in enumerate(items):
            y = 455 + row_idx * 125
            box = (x, y, x + width, y + 78)
            rounded_box(draw, box, item, (247, 248, 250), small_font)
            arrow(draw, (center_x, prev_bottom), (center_x, y))
            prev_bottom = y + 78

    img.save(path)


def crop_without_title(source: Image.Image, top: int, pad: int = 20) -> Image.Image:
    rgb = source.convert("RGB")
    crop = rgb.crop((0, top, rgb.width, rgb.height))
    px = crop.load()
    min_x, min_y, max_x, max_y = crop.width, crop.height, -1, -1
    for y in range(crop.height):
        for x in range(crop.width):
            r, g, b = px[x, y]
            if r < 245 or g < 245 or b < 245:
                min_x = min(min_x, x)
                min_y = min(min_y, y)
                max_x = max(max_x, x)
                max_y = max(max_y, y)
    if max_x < 0:
        return crop
    return crop.crop((max(min_x - pad, 0), max(min_y - pad, 0), min(max_x + pad, crop.width), min(max_y + pad, crop.height)))


def read_media(docx: Path, media_name: str) -> Image.Image:
    with zipfile.ZipFile(docx, "r") as z:
        data = z.read(f"word/media/{media_name}")
    image = Image.open(io.BytesIO(data))
    image.load()
    return image


def image_bytes(image: Image.Image) -> bytes:
    out = io.BytesIO()
    image.save(out, format="PNG")
    return out.getvalue()


def set_paragraph_text(paragraph, text: str) -> None:
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)


def update_text_stage() -> None:
    doc = Document(str(INPUT_DOCX))
    set_paragraph_text(
        doc.paragraphs[258],
        "Навигационная модель начинается с общей карты экранов: она показывает переход от входа и авторизации к разделам, доступным пользователю в зависимости от его роли.",
    )
    set_paragraph_text(
        doc.paragraphs[261],
        "На рисунках 22.1-22.6 приведены общая карта экранов и детальные карты по ролям. Стрелки показывают основные переходы между разделами интерфейса и помогают сопоставить навигацию с правами пользователя.",
    )
    set_paragraph_text(doc.paragraphs[263], "Рисунок 22.1 - Общая карта экранов информационной системы Пульс")

    doc.save(TEXT_STAGE)


def rel_media_for_paragraph(doc: Document, paragraph_index: int) -> str:
    rels = doc.part.rels
    blips = doc.paragraphs[paragraph_index]._p.xpath(".//a:blip")
    if not blips:
        raise RuntimeError(f"paragraph {paragraph_index} has no image")
    rid = blips[0].get(qn("r:embed"))
    target = rels[rid].target_ref
    return f"word/{target}"


def extent_for_paragraph(doc: Document, paragraph_index: int) -> tuple[str, int]:
    blips = doc.paragraphs[paragraph_index]._p.xpath(".//a:blip")
    rid = blips[0].get(qn("r:embed"))
    extent = doc.paragraphs[paragraph_index]._p.xpath(".//wp:extent")[0]
    cx = int(extent.get("cx"))
    return rid, cx


def set_paragraph_image_extent(doc: Document, paragraph_index: int, image_size: tuple[int, int]) -> None:
    paragraph = doc.paragraphs[paragraph_index]
    extents = paragraph._p.xpath(".//wp:extent")
    a_extents = paragraph._p.xpath(".//a:xfrm/a:ext")
    if not extents:
        raise RuntimeError(f"paragraph {paragraph_index} has no extent")
    width_px, height_px = image_size
    cx = int(extents[0].get("cx"))
    cy = int(cx * height_px / width_px)
    extents[0].set("cy", str(cy))
    if a_extents:
        a_extents[0].set("cy", str(cy))


def replace_media_only(source_docx: Path, replacements: dict[str, tuple[bytes, tuple[int, int]]]) -> None:
    with zipfile.ZipFile(source_docx, "r") as zin:
        entries = {name: zin.read(name) for name in zin.namelist()}

    with zipfile.ZipFile(OUTPUT_DOCX, "w", compression=zipfile.ZIP_DEFLATED) as zout:
        for name, data in entries.items():
            if name in replacements:
                data = replacements[name][0]
            zout.writestr(name, data)


def main() -> None:
    if WORK_DIR.exists():
        shutil.rmtree(WORK_DIR)
    WORK_DIR.mkdir(parents=True, exist_ok=True)

    create_general_screen_map(WORK_DIR / "general_screen_map.png")
    update_text_stage()

    stage_doc = Document(str(TEXT_STAGE))
    fig2_media = rel_media_for_paragraph(stage_doc, 40)
    fig9_media = rel_media_for_paragraph(stage_doc, 83)
    general_map_media = rel_media_for_paragraph(stage_doc, 262)

    fig2 = crop_without_title(read_media(REFERENCE_DOCX, "image2.png"), top=70, pad=22)
    fig9 = crop_without_title(read_media(INPUT_DOCX, "image5.png"), top=58, pad=24)
    general = Image.open(WORK_DIR / "general_screen_map.png")
    general.load()

    replacements = {
        fig2_media: (image_bytes(fig2), fig2.size),
        fig9_media: (image_bytes(fig9), fig9.size),
        general_map_media: (image_bytes(general), general.size),
    }

    set_paragraph_image_extent(stage_doc, 40, fig2.size)
    set_paragraph_image_extent(stage_doc, 83, fig9.size)
    set_paragraph_image_extent(stage_doc, 262, general.size)
    stage_doc.save(TEXT_STAGE)
    replace_media_only(TEXT_STAGE, replacements)

    print(f"created={OUTPUT_DOCX}")
    print(f"fig2_media={fig2_media} size={fig2.size}")
    print(f"fig9_media={fig9_media} size={fig9.size}")
    print(f"general_map_media={general_map_media} size={general.size}")


if __name__ == "__main__":
    main()
