from __future__ import annotations

import re
import shutil
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from docx.text.paragraph import Paragraph


FINAL_DIR = Path(r"E:\diploma\docs\diploma_document\zozin_final_diploma_files")
INPUT_DOCX = FINAL_DIR / "Zozin_v9.docx"
OUTPUT_DOCX = FINAL_DIR / "Zozin_v10.docx"
SQL_PATH = Path(r"E:\diploma\unityaid-schema-only.sql")


TYPE_PURPOSES = {
    "uuid": "Идентификатор записи или ссылка на связанную сущность.",
    "text": "Текстовое значение, используемое в интерфейсе или бизнес-логике.",
    "integer": "Целочисленный показатель, статусный параметр или порядок сортировки.",
    "bigint": "Служебный числовой идентификатор для отображения и последовательностей.",
    "boolean": "Логический признак включения, активности или выполнения условия.",
    "numeric": "Числовое значение с фиксированной точностью.",
    "timestamp with time zone": "Дата и время события с учетом часового пояса.",
    "jsonb": "Структурированные данные произвольной формы.",
    "inet": "IP-адрес клиента или сетевого события.",
}

TABLE_PURPOSES = {
    "achievement_rules": "Правила автоматического начисления достижений.",
    "achievements": "Справочник достижений и наград системы.",
    "audit_log": "Журнал значимых пользовательских и системных действий.",
    "certificates": "Сведения о сформированных сертификатах и проверочных кодах.",
    "email_verification_tokens": "Токены подтверждения адреса электронной почты.",
    "event_applications": "Заявки пользователей на участие в мероприятиях.",
    "event_attendance": "Факты посещения мероприятий и рассчитанные часы.",
    "event_feedback": "Отзывы участников о мероприятиях.",
    "event_shifts": "Смены и временные интервалы внутри мероприятий.",
    "event_templates": "Шаблоны мероприятий организации.",
    "events": "Мероприятия, публикуемые организацией.",
    "field_checkins": "Полевые отметки прихода и ухода по QR-маршруту.",
    "field_qr_tokens": "QR-токены для регистрации на месте проведения.",
    "knowledge_articles": "Статьи базы знаний.",
    "knowledge_categories": "Разделы базы знаний.",
    "news": "Новостные публикации.",
    "news_categories": "Категории новостей.",
    "notifications": "Персональные уведомления пользователей.",
    "organization_members": "Членство пользователей в организациях и роли.",
    "organizations": "Организации, работающие в системе.",
    "password_reset_tokens": "Токены восстановления пароля.",
    "points_transactions": "История начисления и списания баллов.",
    "profile_field_definitions": "Настраиваемые поля профиля участника.",
    "profile_field_groups": "Группы настраиваемых полей профиля.",
    "profile_field_options": "Варианты значений для справочных полей профиля.",
    "profile_field_values": "Значения настраиваемых полей пользователей.",
    "refresh_sessions": "Сессии обновления токенов авторизации.",
    "revoked_access_tokens": "Отозванные access-токены.",
    "schema_migrations": "История примененных миграций схемы.",
    "seed_migrations": "История примененных миграций начальных данных.",
    "skills": "Справочник навыков волонтеров.",
    "system_roles": "Справочник системных ролей.",
    "task_assignments": "Назначения пользователей на задачи.",
    "task_attachments": "Файлы, приложенные к задачам.",
    "task_comments": "Комментарии к задачам.",
    "task_status_history": "История изменения статусов задач.",
    "task_time_entries": "Учет времени по задачам.",
    "tasks": "Задачи организации и мероприятий.",
    "tenant_settings": "Глобальные настройки экземпляра системы.",
    "time_entries": "Записи волонтерского времени.",
    "user_invitations": "Приглашения пользователей в организацию.",
    "user_system_roles": "Связь пользователей с системными ролями.",
    "users": "Учетные записи пользователей.",
    "volunteer_achievements": "Факты получения достижений пользователями.",
    "volunteer_profiles": "Профили волонтеров и агрегированные показатели.",
    "volunteer_skills": "Связь профилей волонтеров с навыками.",
}


def set_text(paragraph, text: str) -> None:
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)


def replace_in_paragraph(paragraph, replacements: dict[str, str]) -> None:
    text = paragraph.text
    new = text
    for old, repl in replacements.items():
        new = new.replace(old, repl)
    if new != text:
        set_text(paragraph, new)


def delete_paragraph(paragraph) -> None:
    p = paragraph._element
    p.getparent().remove(p)


def add_table_borders(table) -> None:
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), "000000")


def format_table(table) -> None:
    add_table_borders(table)
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.name = "Times New Roman"
                    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
                    run.font.size = Pt(9)
    for cell in table.rows[0].cells:
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.bold = True


def parse_schema() -> list[dict]:
    sql = SQL_PATH.read_text(encoding="utf-8", errors="ignore")
    table_re = re.compile(r"CREATE TABLE public\.([A-Za-z_][\w]*) \((.*?)\n\);", re.S)
    constraints: dict[str, list[str]] = {}
    for match in re.finditer(r"ALTER TABLE ONLY public\.([A-Za-z_][\w]*)\s+ADD CONSTRAINT\s+([A-Za-z_][\w]*)\s+(.+?);", sql, re.S):
        table, name, body = match.groups()
        body = " ".join(body.split())
        constraints.setdefault(table, []).append(body)

    foreign_keys: dict[str, dict[str, str]] = {}
    for table, cons in constraints.items():
        for body in cons:
            fk = re.search(r"FOREIGN KEY \(([^)]+)\) REFERENCES public\.([A-Za-z_][\w]*)\(([^)]+)\)", body)
            if fk:
                cols = [c.strip() for c in fk.group(1).split(",")]
                ref = f"{fk.group(2)}({fk.group(3)})"
                for col in cols:
                    foreign_keys.setdefault(table, {})[col] = ref

    tables = []
    for table_name, body in table_re.findall(sql):
        table_constraints = constraints.get(table_name, [])
        columns = []
        for raw in body.splitlines():
            line = raw.strip().rstrip(",")
            if not line or line.startswith(("CONSTRAINT", "PRIMARY KEY", "FOREIGN KEY", "UNIQUE", "CHECK")):
                continue
            m = re.match(r"([A-Za-z_][\w]*)\s+(.+)", line)
            if not m:
                continue
            col, type_and_rules = m.groups()
            col_type = type_and_rules
            for token in (" DEFAULT ", " NOT NULL", " COLLATE "):
                if token in col_type:
                    col_type = col_type.split(token)[0]
            col_type = col_type.replace("public.", "").strip()
            rules = []
            if "NOT NULL" in type_and_rules:
                rules.append("NOT NULL")
            if "DEFAULT" in type_and_rules:
                rules.append("DEFAULT")
            for cons in table_constraints:
                if f"PRIMARY KEY ({col})" in cons:
                    rules.append("PK")
                if re.search(rf"UNIQUE \([^)]*\b{re.escape(col)}\b[^)]*\)", cons):
                    rules.append("UNIQUE")
            if col in foreign_keys.get(table_name, {}):
                rules.append("FK -> " + foreign_keys[table_name][col])
            columns.append({"name": col, "type": col_type, "rules": "; ".join(dict.fromkeys(rules)) or "-", "purpose": field_purpose(col, col_type)})
        tables.append({"name": table_name, "columns": columns, "purpose": TABLE_PURPOSES.get(table_name, "Таблица базы данных информационной системы Пульс.")})
    return tables


def field_purpose(name: str, col_type: str) -> str:
    if name == "id":
        return "Первичный идентификатор записи."
    if name == "display_id":
        return "Числовой идентификатор для отображения и сортировки."
    if name.endswith("_id"):
        return "Ссылка на связанную сущность."
    if name in {"created_at", "updated_at", "assigned_at", "earned_at", "revoked_at", "used_at", "read_at"}:
        return "Дата и время фиксации соответствующего события."
    if name in {"starts_at", "ends_at", "expires_at", "due_at", "last_login_at"}:
        return "Дата и время, определяющие период или срок действия."
    if name in {"status", "role", "type", "format", "priority", "mode", "source"}:
        return "Код состояния, роли, режима или классификации."
    if name in {"title", "name", "code", "slug", "email", "token", "token_hash", "verify_code"}:
        return "Именующее или уникальное значение для поиска и идентификации."
    if name in {"description", "summary", "body", "content", "content_html", "comment", "message", "reason", "note"}:
        return "Текстовое описание или комментарий."
    if name in {"hours", "total_hours", "points", "points_reward", "level", "rating", "capacity", "max_participants", "sort_order"}:
        return "Количественный показатель, используемый в предметной логике."
    if name.startswith("is_") or name in {"required", "is_active", "is_system", "is_read"}:
        return "Логический признак записи."
    for key, purpose in TYPE_PURPOSES.items():
        if col_type.startswith(key):
            return purpose
    return "Поле таблицы, используемое бизнес-логикой системы."


def add_caption(doc: Document, text: str, template=None) -> None:
    p = doc.add_paragraph()
    if template is not None:
        p.style = template.style
        p.alignment = template.alignment
    else:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(12)


def add_appendix(doc: Document, tables: list[dict]) -> None:
    doc.add_page_break()
    heading = doc.add_paragraph()
    heading.style = doc.paragraphs[687].style
    set_text(heading, "ПРИЛОЖЕНИЕ Д Полное физическое описание таблиц базы данных")

    intro = doc.add_paragraph()
    intro.style = doc.paragraphs[688].style
    set_text(
        intro,
        "В приложении приведено полное физическое описание 46 таблиц базы данных PostgreSQL информационной системы Пульс. Для каждой таблицы указаны поля, типы данных, ключевые ограничения и назначение полей.",
    )

    caption_template = doc.paragraphs[257]
    for idx, table_info in enumerate(tables, start=1):
        add_caption(doc, f"Таблица Д.{idx} - Физическое описание таблицы {table_info['name']}", caption_template)
        table = doc.add_table(rows=1, cols=4)
        table.autofit = True
        headers = ["Поле", "Тип", "Ограничения и связи", "Назначение"]
        for cell, header in zip(table.rows[0].cells, headers):
            cell.text = header
        for col in table_info["columns"]:
            row = table.add_row().cells
            row[0].text = col["name"]
            row[1].text = col["type"]
            row[2].text = col["rules"]
            row[3].text = col["purpose"]
        format_table(table)


def main() -> None:
    shutil.copyfile(INPUT_DOCX, OUTPUT_DOCX)
    doc = Document(str(OUTPUT_DOCX))

    # Minimal fixes on title/task/review pages: content only, styles stay in place.
    replace_in_paragraph(doc.paragraphs[12], {"–": "-"})
    replace_in_paragraph(doc.paragraphs[77], {"информационной системы пульс": "информационной системы Пульс"})
    replace_in_paragraph(doc.paragraphs[99], {"—": "-"})
    replace_in_paragraph(doc.paragraphs[102], {"—": "-", "»..": "»."})

    abstract_text = (
        "Пояснительная записка содержит введение, 3 раздела, заключение, 97 таблиц, 51 рисунок, список из 46 использованных источников и приложения. "
        "В работе представлены модели предметной области и данных, архитектура, алгоритмические схемы, макеты интерфейса, результаты сборки и программа опытной эксплуатации продукта."
    )

    # Remove figure references for screenshots that are not present in the final document.
    set_text(doc.paragraphs[156], "Платформа Добро.рф использована как пример публичной добровольческой экосистемы при сравнении аналогов.")
    set_text(doc.paragraphs[158], "Публичная модель сервиса VolunteerMatch на платформе Idealist учитывалась при оценке сценария поиска волонтерских возможностей.")
    set_text(doc.paragraphs[160], "Сервис Timecounts использован при сравнении как пример решения, ориентированного на операционную работу координатора.")
    set_text(doc.paragraphs[162], "Интерфейс и функциональная модель Golden учитывались как пример коммерческой SaaS-платформы управления волонтерами.")

    replacements = {
        "рисунке 8": "рисунке 4",
        "Рисунок 8 -": "Рисунок 4 -",
        "рисунке 9": "рисунке 5",
        "Рисунок 9 -": "Рисунок 5 -",
        "рисунке 10": "рисунке 6",
        "Рисунок 10 -": "Рисунок 6 -",
        "рисунке 11": "рисунке 7",
        "Рисунок 11 -": "Рисунок 7 -",
        "рисунке 12": "рисунке 8",
        "Рисунок 12 -": "Рисунок 8 -",
        "рисунке 13": "рисунке 9",
        "Рисунок 13 -": "Рисунок 9 -",
        "рисунках 15-18": "рисунках 10-13",
        "Рисунок 15 -": "Рисунок 10 -",
        "Рисунок 16 -": "Рисунок 11 -",
        "Рисунок 17 - ER": "Рисунок 12 - ER",
        "Рисунок 18 - ER": "Рисунок 13 - ER",
        "рисунках 17-21": "рисунках 14-18",
        "Рисунок 17 - Блок": "Рисунок 14 - Блок",
        "Рисунок 18 - Блок": "Рисунок 15 - Блок",
        "Рисунок 19 -": "Рисунок 16 -",
        "Рисунок 20 -": "Рисунок 17 -",
        "Рисунок 21 -": "Рисунок 18 -",
        "22.1-22.6": "19.1-19.6",
        "Рисунок 22.1 -": "Рисунок 19.1 -",
        "Рисунок 22.2 -": "Рисунок 19.2 -",
        "Рисунок 22.3 -": "Рисунок 19.3 -",
        "Рисунок 22.4 -": "Рисунок 19.4 -",
        "Рисунок 22.5 -": "Рисунок 19.5 -",
        "Рисунок 22.6 -": "Рисунок 19.6 -",
        "рисунке 23": "рисунке 20",
        "Рисунок 23 -": "Рисунок 20 -",
        "рисунке 24": "рисунке 21",
        "Рисунок 24 -": "Рисунок 21 -",
        "рисунке 25": "рисунке 22",
        "Рисунок 25 -": "Рисунок 22 -",
        "рисунке 26": "рисунке 23",
        "Рисунок 26 -": "Рисунок 23 -",
        "Рисунок 34 -": "Рисунок 24 -",
        "Рисунок 35 -": "Рисунок 25 -",
        "Рисунок 36 -": "Рисунок 26 -",
        "Рисунок 37 -": "Рисунок 27 -",
        "Рисунок 38 -": "Рисунок 28 -",
        "Рисунок 39 -": "Рисунок 29 -",
        "Рисунок 40 -": "Рисунок 30 -",
        "Рисунок 41 -": "Рисунок 31 -",
        "Рисунок 42 -": "Рисунок 32 -",
        "Рисунок 43 -": "Рисунок 33 -",
        "рисунке 34": "рисунке 24",
        "рисунке 35": "рисунке 25",
        "рисунке 36": "рисунке 26",
        "рисунке 37": "рисунке 27",
        "рисунке 38": "рисунке 28",
        "рисунке 39": "рисунке 29",
        "рисунке 40": "рисунке 30",
        "рисунке 41": "рисунке 31",
        "рисунке 42": "рисунке 32",
        "рисунке 43": "рисунке 33",
        "Таблица 27 - Основные группы API": "Таблица 26 - Основные группы API",
        "Таблица 28 - Реализованные модули Пульс": "Таблица 27 - Реализованные модули Пульс",
        "Таблица 26 - Карта экранов информационной системы Пульс": "Таблица 28 - Карта экранов информационной системы Пульс",
        "таблице 27": "таблице 26",
        "таблице 28": "таблице 27",
        "таблице 33": "таблице 50",
        "таблицах 16-33": "таблицах 33-50",
        "Таблица В.1": "Таблица Б.1",
        "таблице В.1": "таблице Б.1",
        "рисунках Г.1-Г.5": "рисунках В.1-В.5",
        "Рисунок Г.1 - Обработка": "Рисунок В.1 - Обработка",
        "Рисунок Г.2 - Подтверждение": "Рисунок В.2 - Подтверждение",
        "Рисунок Г.3 - Начисление": "Рисунок В.3 - Начисление",
        "Рисунок Г.4 - Формирование сертификата": "Рисунок В.4 - Формирование сертификата",
        "Рисунок Г.5 - Формирование аналитики": "Рисунок В.5 - Формирование аналитики",
        "рисунках Д.1-Д.10": "рисунках Г.1-Г.8",
        "Листинг Д.1": "Листинг Г.1",
        "—": "-",
        "–": "-",
    }
    for paragraph in doc.paragraphs:
        replace_in_paragraph(paragraph, replacements)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    replace_in_paragraph(paragraph, replacements)

    set_text(
        doc.paragraphs[256],
        "Физическое описание в основной части приведено для ключевых таблиц, которые образуют предметное ядро и показаны на ER-диаграммах. Полное физическое описание всех 46 таблиц базы данных PostgreSQL вынесено в приложение Д в конце пояснительной записки.",
    )
    marker = doc.paragraphs[294]
    insert_paragraph_before(
        marker,
        "Перечень таблиц приложения Д включает предметные, справочные, служебные и технические таблицы: сессии, токены, миграции, настройки, профильные поля, QR-отметки, шаблоны мероприятий и расширения задач.",
    )

    delete_paragraph(doc.paragraphs[103])
    supervisor_line = "\u0420\u0443\u043a\u043e\u0432\u043e\u0434\u0438\u0442\u0435\u043b\u044c \u0432\u044b\u043f\u0443\u0441\u043a\u043d\u043e\u0439 \u043a\u0432\u0430\u043b\u0438\u0444\u0438\u043a\u0430\u0446\u0438\u043e\u043d\u043d\u043e\u0439 \u0440\u0430\u0431\u043e\u0442\u044b,"
    signature_prefix = "\u043a.\u043f.\u043d., \u0434\u043e\u0446\u0435\u043d\u0442 \u043a\u0430\u0444\u0435\u0434\u0440\u044b"
    for idx, paragraph in enumerate(doc.paragraphs):
        if paragraph.text.strip().startswith(signature_prefix):
            previous_text = ""
            for previous in reversed(doc.paragraphs[:idx]):
                if previous.text.strip():
                    previous_text = previous.text.strip()
                    break
            if previous_text != supervisor_line:
                insert_paragraph_before(paragraph, supervisor_line)
            break

    abstract_seen = False
    for paragraph in list(doc.paragraphs):
        text = paragraph.text.strip()
        if text.startswith("Пояснительная записка содержит введение"):
            if abstract_seen:
                delete_paragraph(paragraph)
            else:
                set_text(paragraph, abstract_text)
                abstract_seen = True
        elif text.startswith("Практическая значимость результата") and not text.endswith("."):
            set_text(paragraph, text + ".")
        elif ". подтверждаемая" in text:
            replace_in_paragraph(paragraph, {". подтверждаемая": ". Подтверждаемая"})

    add_appendix(doc, parse_schema())
    doc.save(OUTPUT_DOCX)
    print(OUTPUT_DOCX)


def insert_paragraph_before(paragraph, text: str):
    new_p = OxmlElement("w:p")
    paragraph._p.addprevious(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    new_para.style = paragraph.style
    set_text(new_para, text)
    return new_para


if __name__ == "__main__":
    main()
