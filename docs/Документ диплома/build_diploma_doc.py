from __future__ import annotations

from pathlib import Path
from typing import Iterable

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[2]
OUT_DIR = ROOT / "docs" / "Документ диплома" / "generated"
IMG_DIR = OUT_DIR / "images"
DOCX_PATH = OUT_DIR / "Zozin_UnityAid_VKR_final.docx"
EXAMPLE_TEMPLATE = OUT_DIR.parent / "Примеры" / "Diplom_Kustubaev_AVb-22-2.docx"

TITLE = "Проектирование и разработка информационной системы управления  волонтерами с элементами геймификации и аналитики"
STUDENT = "Зозин Павел Алексеевич"
GROUP = "Авб-22-2"
DEPARTMENT = "вычислительной техники и программирования"
SUPERVISOR = "Гладышева Мария Михайловна"
YEAR = "2026"


def set_cell_text(cell, text: str, bold: bool = False, align=WD_ALIGN_PARAGRAPH.LEFT, font_size: int = 12):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    parts = str(text).split("\n")
    for idx, part in enumerate(parts):
        if idx:
            p.add_run().add_break()
        run = p.add_run(part)
        run.bold = bold
        run.font.name = "Times New Roman"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        run.font.size = Pt(font_size)
    for paragraph in cell.paragraphs:
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.0


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr)
    run._r.append(fld_char2)


def add_toc(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.first_line_indent = Cm(0)
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = r'TOC \o "1-3" \h \z \u'
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "Оглавление обновляется автоматически"
    fld_char3 = OxmlElement("w:fldChar")
    fld_char3.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr)
    run._r.append(fld_char2)
    run._r.append(placeholder)
    run._r.append(fld_char3)


def set_page_number_start(section, start: int):
    sect_pr = section._sectPr
    pg_num = sect_pr.find(qn("w:pgNumType"))
    if pg_num is None:
        pg_num = OxmlElement("w:pgNumType")
        sect_pr.append(pg_num)
    pg_num.set(qn("w:start"), str(start))


def clear_header_footer(part):
    element = part._element
    for child in list(element):
        element.remove(child)
    return part.add_paragraph()


def configure_section(section, page_numbers: bool = False, start: int | None = None):
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.left_margin = Mm(30)
    section.right_margin = Mm(10)
    section.top_margin = Mm(20)
    section.bottom_margin = Mm(20)
    section.header_distance = Mm(12)
    section.footer_distance = Mm(12)
    section.footer.is_linked_to_previous = False
    section.header.is_linked_to_previous = False
    clear_header_footer(section.header)
    footer_paragraph = clear_header_footer(section.footer)
    if page_numbers:
        if start is not None:
            set_page_number_start(section, start)
        add_page_number(footer_paragraph)


def clear_document_body(doc: Document):
    body = doc._body._element
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def set_styles(doc: Document):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(14)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    pf = normal.paragraph_format
    pf.first_line_indent = Cm(1.25)
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf.line_spacing = 1.5
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)

    for name, size, all_caps in [
        ("Heading 1", 14, True),
        ("Heading 2", 14, False),
        ("Heading 3", 14, False),
    ]:
        st = styles[name]
        st.font.name = "Times New Roman"
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor(0, 0, 0)
        st.font.all_caps = all_caps
        st.paragraph_format.first_line_indent = Cm(1.25)
        st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        st.paragraph_format.line_spacing = 1.5
        st.paragraph_format.space_before = Pt(0)
        st.paragraph_format.space_after = Pt(0)
        if name == "Heading 1":
            st.paragraph_format.page_break_before = True

    for style_name in ["Caption", "List Bullet", "List Number"]:
        if style_name not in styles:
            st = styles.add_style(style_name, WD_STYLE_TYPE.PARAGRAPH)
            st.base_style = normal
        else:
            st = styles[style_name]
        st.font.name = "Times New Roman"
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        st.font.size = Pt(14)
        st.font.color.rgb = RGBColor(0, 0, 0)
        st.paragraph_format.line_spacing = 1.5
        st.paragraph_format.space_after = Pt(0)

    if "Table Text" not in styles:
        st = styles.add_style("Table Text", WD_STYLE_TYPE.PARAGRAPH)
        st.base_style = normal
    styles["Table Text"].font.name = "Times New Roman"
    styles["Table Text"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    styles["Table Text"].font.size = Pt(12)
    styles["Table Text"].paragraph_format.first_line_indent = Cm(0)
    styles["Table Text"].paragraph_format.line_spacing = 1.0


def paragraph(doc: Document, text: str = "", style: str | None = None, align=None, first_indent=True, bold=False):
    p = doc.add_paragraph(style=style)
    if align is not None:
        p.alignment = align
    if not first_indent:
        p.paragraph_format.first_line_indent = Cm(0)
    run = p.add_run(text)
    run.bold = bold
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    return p


def heading(doc: Document, text: str, level: int):
    p = doc.add_heading(text, level=level)
    if level == 1:
        centered = not text[:1].isdigit()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if centered else WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.first_line_indent = Cm(0 if centered else 1.25)
    return p


def bullet(doc: Document, items: Iterable[str]):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Cm(1.25)
        p.paragraph_format.first_line_indent = Cm(0)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.add_run(item)


def number_list(doc: Document, items: Iterable[str]):
    for idx, item in enumerate(items, 1):
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Cm(1.25)
        p.paragraph_format.left_indent = Cm(0)
        p.paragraph_format.line_spacing = 1.5
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        run = p.add_run(f"{idx}) {item}")
        run.font.name = "Times New Roman"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        run.font.size = Pt(14)


def set_table_fixed_geometry(tbl, widths: list[float]):
    tbl.autofit = False
    tbl_pr = tbl._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    total_twips = sum(int(w * 567) for w in widths)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(total_twips))
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    tbl_grid = tbl._tbl.tblGrid
    if tbl_grid is None:
        tbl_grid = OxmlElement("w:tblGrid")
        tbl._tbl.insert(0, tbl_grid)
    for child in list(tbl_grid):
        tbl_grid.remove(child)
    for width in widths:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(int(width * 567)))
        tbl_grid.append(grid_col)
    for row in tbl.rows:
        for idx, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:type"), "dxa")
            tc_w.set(qn("w:w"), str(int(widths[idx] * 567)))


def table(doc: Document, caption: str, headers: list[str], rows: list[list[str]], widths: list[float], font_size: int = 12):
    paragraph(doc, caption, align=WD_ALIGN_PARAGRAPH.LEFT, first_indent=False)
    tbl = doc.add_table(rows=1, cols=len(headers))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.style = "Table Grid"
    set_table_fixed_geometry(tbl, widths)
    for idx, h in enumerate(headers):
        cell = tbl.rows[0].cells[idx]
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_text(cell, h, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, font_size=font_size)
        cell.width = Cm(widths[idx])
    for row in rows:
        cells = tbl.add_row().cells
        for idx, value in enumerate(row):
            cells[idx].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_text(cells[idx], value, align=WD_ALIGN_PARAGRAPH.LEFT, font_size=font_size)
            cells[idx].width = Cm(widths[idx])
    set_table_fixed_geometry(tbl, widths)
    paragraph(doc, "")
    return tbl


def draw_box_diagram(path: Path, title: str, boxes: list[tuple[int, int, int, int, str]], arrows: list[tuple[int, int, int, int]]):
    img = Image.new("RGB", (1400, 850), "white")
    draw = ImageDraw.Draw(img)
    try:
        font_title = ImageFont.truetype("arial.ttf", 34)
        font = ImageFont.truetype("arial.ttf", 24)
        font_small = ImageFont.truetype("arial.ttf", 20)
    except Exception:
        font_title = font = font_small = ImageFont.load_default()
    draw.text((40, 30), title, fill=(0, 0, 0), font=font_title)
    for x1, y1, x2, y2, label in boxes:
        draw.rounded_rectangle((x1, y1, x2, y2), radius=18, outline=(0, 0, 0), width=3, fill=(245, 247, 250))
        lines = label.split("\n")
        total_h = len(lines) * 30
        start_y = y1 + ((y2 - y1) - total_h) / 2
        for line in lines:
            bbox = draw.textbbox((0, 0), line, font=font_small if len(line) > 32 else font)
            draw.text((x1 + ((x2 - x1) - (bbox[2] - bbox[0])) / 2, start_y), line, fill=(0, 0, 0), font=font_small if len(line) > 32 else font)
            start_y += 30
    for x1, y1, x2, y2 in arrows:
        draw.line((x1, y1, x2, y2), fill=(0, 0, 0), width=3)
        ang = 12
        if x2 >= x1:
            draw.polygon([(x2, y2), (x2 - 20, y2 - ang), (x2 - 20, y2 + ang)], fill=(0, 0, 0))
        else:
            draw.polygon([(x2, y2), (x2 + 20, y2 - ang), (x2 + 20, y2 + ang)], fill=(0, 0, 0))
    img.save(path)


def make_images():
    IMG_DIR.mkdir(parents=True, exist_ok=True)
    draw_box_diagram(
        IMG_DIR / "architecture.png",
        "Архитектура UnityAid",
        [
            (60, 160, 360, 300, "Пользователь\nбраузер / PWA"),
            (520, 130, 880, 330, "Frontend\nVue 3, TypeScript, Vite"),
            (1040, 130, 1320, 330, "Backend API\nGo, REST"),
            (1040, 450, 1320, 610, "PostgreSQL\nданные системы"),
            (520, 450, 880, 610, "Worker\nуведомления, отчеты"),
            (60, 450, 360, 610, "Файловое хранилище\nаватары, сертификаты"),
        ],
        [(360, 230, 520, 230), (880, 230, 1040, 230), (1180, 330, 1180, 450), (1040, 530, 880, 530), (520, 530, 360, 530)],
    )
    draw_box_diagram(
        IMG_DIR / "process.png",
        "Сценарий проведения мероприятия",
        [
            (60, 170, 300, 300, "Координатор\nсоздает мероприятие"),
            (380, 170, 620, 300, "Волонтер\nподает заявку"),
            (700, 170, 940, 300, "Координатор\nподтверждает"),
            (1020, 170, 1300, 300, "Проведение\nи отметка участия"),
            (380, 460, 620, 590, "Учет часов"),
            (700, 460, 940, 590, "Начисление\nбаллов"),
            (1020, 460, 1300, 590, "Сертификат\nи отчет"),
        ],
        [(300, 235, 380, 235), (620, 235, 700, 235), (940, 235, 1020, 235), (1160, 300, 1160, 460), (1020, 525, 940, 525), (700, 525, 620, 525)],
    )
    draw_box_diagram(
        IMG_DIR / "modules.png",
        "Функциональные модули системы",
        [
            (520, 340, 880, 500, "UnityAid\nядро приложения"),
            (60, 120, 350, 250, "Пользователи\nи роли"),
            (540, 120, 860, 250, "Организации\nи участники"),
            (1050, 120, 1320, 250, "Мероприятия\nи заявки"),
            (60, 580, 350, 710, "Задачи\nи назначения"),
            (540, 580, 860, 710, "Часы\nи сертификаты"),
            (1050, 580, 1320, 710, "Аналитика\nи аудит"),
        ],
        [(350, 185, 520, 390), (700, 250, 700, 340), (1050, 185, 880, 390), (350, 645, 520, 450), (700, 580, 700, 500), (1050, 645, 880, 450)],
    )


def add_picture(doc: Document, image: Path, caption: str):
    p = paragraph(doc, "", align=WD_ALIGN_PARAGRAPH.CENTER, first_indent=False)
    p.add_run().add_picture(str(image), width=Cm(15.5))
    cap = paragraph(doc, caption, style="Caption", align=WD_ALIGN_PARAGRAPH.CENTER, first_indent=False)
    cap.runs[0].italic = False


def add_screenshot(doc: Document, image: Path, caption: str):
    if not image.exists():
        return
    p = paragraph(doc, "", align=WD_ALIGN_PARAGRAPH.CENTER, first_indent=False)
    p.add_run().add_picture(str(image), width=Cm(15.0))
    cap = paragraph(doc, caption, style="Caption", align=WD_ALIGN_PARAGRAPH.CENTER, first_indent=False)
    cap.runs[0].italic = False


def code_listing(doc: Document, caption: str, source: Path, start: str | None = None, end: str | None = None, max_lines: int = 80):
    paragraph(doc, caption, align=WD_ALIGN_PARAGRAPH.LEFT, first_indent=False)
    if not source.exists():
        paragraph(doc, f"Файл {source.name} не найден.", first_indent=False)
        return
    lines = source.read_text(encoding="utf-8").splitlines()
    if start:
        for idx, line in enumerate(lines):
            if start in line:
                lines = lines[idx:]
                break
    if end:
        for idx, line in enumerate(lines[1:], 1):
            if end in line:
                lines = lines[:idx]
                break
    lines = lines[:max_lines]
    for raw in lines:
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.left_indent = Cm(0)
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(raw[:135])
        run.font.name = "Courier New"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Courier New")
        run.font.size = Pt(8)
    paragraph(doc, "")


def add_screenshots_appendix(doc: Document):
    heading(doc, "ПРИЛОЖЕНИЕ Д Экранные формы приложения", 1)
    paragraph(doc, "В приложении приведены основные экранные формы разработанной информационной системы UnityAid. Скриншоты подтверждают наличие пользовательского интерфейса для авторизации, работы с мероприятиями, задачами, учетом часов, аналитикой, сертификатами и системным администрированием.", first_indent=False)
    screenshots = [
        ("01_login.png", "Рисунок Д.1 - Страница авторизации пользователя"),
        ("02_dashboard.png", "Рисунок Д.2 - Главная панель пользователя"),
        ("03_events.png", "Рисунок Д.3 - Раздел управления мероприятиями"),
        ("04_tasks.png", "Рисунок Д.4 - Раздел задач"),
        ("05_time_entries.png", "Рисунок Д.5 - Раздел учета волонтерских часов"),
        ("06_analytics.png", "Рисунок Д.6 - Аналитическая панель"),
        ("07_certificates.png", "Рисунок Д.7 - Раздел сертификатов"),
        ("08_admin.png", "Рисунок Д.8 - Системная админ-панель"),
    ]
    shot_dir = IMG_DIR / "screenshots"
    for name, caption in screenshots:
        add_screenshot(doc, shot_dir / name, caption)


def add_code_appendix(doc: Document):
    heading(doc, "ПРИЛОЖЕНИЕ Е Фрагменты исходного кода", 1)
    paragraph(doc, "В приложении приведены фрагменты исходного кода, отражающие ключевые части реализации UnityAid: применение миграций базы данных, работу административного API, маршрутизацию клиентского приложения, взаимодействие frontend с REST API и реализацию бизнес-операций.", first_indent=False)
    code_listing(
        doc,
        "Листинг Е.1 - Миграция для пользовательской нумерации записей",
        ROOT / "unityaid-back" / "migrations" / "019_display_ids.sql",
        max_lines=70,
    )
    code_listing(
        doc,
        "Листинг Е.2 - Формирование списка записей в административном модуле backend",
        ROOT / "unityaid-back" / "internal" / "modules" / "admin" / "repository.go",
        start="func (r *Repository) List",
        end="func (r *Repository) Create",
        max_lines=65,
    )
    code_listing(
        doc,
        "Листинг Е.3 - Клиентские функции административного REST API",
        ROOT / "unityaid-front" / "src" / "entities" / "admin" / "api.ts",
        max_lines=60,
    )
    code_listing(
        doc,
        "Листинг Е.4 - Регистрация маршрутов клиентского приложения",
        ROOT / "unityaid-front" / "src" / "router" / "index.ts",
        max_lines=85,
    )
    code_listing(
        doc,
        "Листинг Е.5 - Обработка заявок на мероприятия",
        ROOT / "unityaid-back" / "internal" / "modules" / "events" / "operations.go",
        start="func (r *Repository) CreateApplication",
        end="func (r *Repository) UpdateApplicationStatus",
        max_lines=90,
    )
    code_listing(
        doc,
        "Листинг Е.6 - Сбор аналитических показателей",
        ROOT / "unityaid-back" / "internal" / "modules" / "analytics" / "repository.go",
        start="func (r *Repository) Overview",
        max_lines=75,
    )


def normalize_black_text(doc: Document):
    for style_name in ["Normal", "Heading 1", "Heading 2", "Heading 3", "Caption", "List Bullet", "List Number"]:
        if style_name in doc.styles:
            doc.styles[style_name].font.color.rgb = RGBColor(0, 0, 0)
    for paragraph_item in doc.paragraphs:
        for run in paragraph_item.runs:
            run.font.color.rgb = RGBColor(0, 0, 0)
    for table_item in doc.tables:
        for row in table_item.rows:
            for cell in row.cells:
                for paragraph_item in cell.paragraphs:
                    for run in paragraph_item.runs:
                        run.font.color.rgb = RGBColor(0, 0, 0)


def title_page(doc: Document):
    for text in [
        "Министерство науки и высшего образования Российской Федерации",
        "Федеральное государственное бюджетное образовательное учреждение",
        "высшего образования",
        "«Магнитогорский государственный технический университет",
        "им. Г.И. Носова»",
        "(ФГБОУ ВО «МГТУ им. Г.И. Носова»)",
    ]:
        p = paragraph(doc, text, align=WD_ALIGN_PARAGRAPH.CENTER, first_indent=False)
        p.paragraph_format.line_spacing = 1.0
    paragraph(doc, "")
    paragraph(doc, f"Кафедра {DEPARTMENT}", align=WD_ALIGN_PARAGRAPH.CENTER, first_indent=False)
    for _ in range(4):
        paragraph(doc, "")
    paragraph(doc, "ВЫПУСКНАЯ КВАЛИФИКАЦИОННАЯ РАБОТА", align=WD_ALIGN_PARAGRAPH.CENTER, first_indent=False, bold=True)
    paragraph(doc, "на тему:", align=WD_ALIGN_PARAGRAPH.CENTER, first_indent=False)
    paragraph(doc, TITLE, align=WD_ALIGN_PARAGRAPH.CENTER, first_indent=False, bold=True)
    paragraph(doc, "по направлению подготовки 09.03.01 Информатика и вычислительная техника", align=WD_ALIGN_PARAGRAPH.CENTER, first_indent=False)
    paragraph(doc, "профиль «Проектирование и разработка Web-приложений»", align=WD_ALIGN_PARAGRAPH.CENTER, first_indent=False)
    for _ in range(3):
        paragraph(doc, "")
    t = doc.add_table(rows=2, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.RIGHT
    for row in t.rows:
        row.cells[0].width = Cm(6)
        row.cells[1].width = Cm(8)
    set_cell_text(t.cell(0, 0), "Исполнитель:", align=WD_ALIGN_PARAGRAPH.RIGHT)
    set_cell_text(t.cell(0, 1), f"{STUDENT}, студент 4 курса, группы {GROUP}")
    set_cell_text(t.cell(1, 0), "Руководитель:", align=WD_ALIGN_PARAGRAPH.RIGHT)
    set_cell_text(t.cell(1, 1), SUPERVISOR)
    for _ in range(7):
        paragraph(doc, "")
    paragraph(doc, f"Магнитогорск, {YEAR}", align=WD_ALIGN_PARAGRAPH.CENTER, first_indent=False)


def assignment_page(doc: Document):
    doc.add_page_break()
    paragraph(doc, "ЗАДАНИЕ", align=WD_ALIGN_PARAGRAPH.CENTER, first_indent=False, bold=True)
    paragraph(doc, "на выполнение выпускной квалификационной работы", align=WD_ALIGN_PARAGRAPH.CENTER, first_indent=False)
    paragraph(doc, f"Обучающемуся {STUDENT}, группа {GROUP}.", first_indent=False)
    paragraph(doc, f"Тема работы: {TITLE}.")
    paragraph(doc, "Исходные данные к работе: материалы технического задания, результаты преддипломной практики, программная реализация UnityAid, проектная документация, нормативные требования к оформлению ВКР.")
    paragraph(doc, "Содержание расчетно-пояснительной записки:")
    number_list(doc, [
        "проанализировать предметную область управления волонтерской деятельностью;",
        "сформулировать требования к информационной системе и определить состав пользовательских ролей;",
        "спроектировать архитектуру web-приложения, структуру базы данных и основные сценарии работы;",
        "описать реализацию клиентской и серверной частей системы UnityAid;",
        "провести проверку работоспособности и оценить результаты разработки.",
    ])
    paragraph(doc, "Перечень графического материала: архитектура системы, схема пользовательского процесса, модель функциональных модулей, структура данных.")
    paragraph(doc, "Дата выдачи задания: «22» апреля 2026 г.")
    paragraph(doc, "Срок представления работы к защите: «___» июня 2026 г.")
    for _ in range(3):
        paragraph(doc, "")
    paragraph(doc, "Руководитель ВКР ____________________ /М.М. Гладышева/", first_indent=False)
    paragraph(doc, "Обучающийся ____________________ /П.А. Зозин/", first_indent=False)


def abstract_page(doc: Document):
    doc.add_section(WD_SECTION.NEW_PAGE)
    configure_section(doc.sections[-1], page_numbers=True, start=3)
    paragraph(doc, "РЕФЕРАТ", align=WD_ALIGN_PARAGRAPH.CENTER, first_indent=False, bold=True)
    paragraph(doc, "Выпускная квалификационная работа содержит 3 раздела, введение, заключение, список использованных источников и приложения. В работе представлены таблицы с функциональными требованиями, структурой базы данных и результатами проверки, а также иллюстрации архитектуры и пользовательских процессов.", first_indent=False)
    paragraph(doc, "Ключевые слова: ИНФОРМАЦИОННАЯ СИСТЕМА, ВОЛОНТЕРЫ, WEB-ПРИЛОЖЕНИЕ, ГЕЙМИФИКАЦИЯ, АНАЛИТИКА, GO, VUE.JS, POSTGRESQL, REST API, СЕРТИФИКАТЫ", first_indent=False)
    paragraph(doc, "Объектом разработки является процесс управления волонтерской деятельностью в организациях, которые проводят мероприятия, распределяют задачи, учитывают вклад участников и формируют отчетность.")
    paragraph(doc, "Цель работы - разработка web-ориентированной информационной системы UnityAid, позволяющей вести учет волонтеров, мероприятий, задач и подтвержденных часов, а также повышать вовлеченность участников за счет достижений, рейтингов и персональной статистики.")
    paragraph(doc, "В ходе работы выполнен анализ предметной области, сформулированы функциональные и нефункциональные требования, спроектирована архитектура приложения, разработаны серверная часть на языке Go, клиентская часть на Vue.js и база данных PostgreSQL. Реализованы модули авторизации, организаций, мероприятий, задач, учета времени, уведомлений, базы знаний, геймификации, аналитики, сертификатов и системного администрирования.")
    paragraph(doc, "Практическая значимость результата заключается в возможности использовать систему как основу для автоматизации деятельности малых и средних волонтерских организаций, студенческих объединений и социальных проектов.")


def contents_page(doc: Document, add_break: bool = True):
    if add_break:
        doc.add_page_break()
    paragraph(doc, "СОДЕРЖАНИЕ", align=WD_ALIGN_PARAGRAPH.CENTER, first_indent=False, bold=True)
    add_toc(doc.add_paragraph())


def introduction(doc: Document):
    heading(doc, "ВВЕДЕНИЕ", 1)
    texts = [
        "Развитие цифровых сервисов изменило подход к организации общественных, образовательных и социальных инициатив. Волонтерские объединения все чаще работают не как разовые группы помощи, а как устойчивые организационные структуры, которым требуется учет участников, планирование мероприятий, распределение задач, подтверждение часов и подготовка отчетов. При отсутствии единой информационной системы эти процессы распределяются между таблицами, мессенджерами, электронными формами и устными договоренностями.",
        "Такая организация работы создает ряд практических проблем. Данные о волонтерах быстро устаревают, история участия хранится неполно, координатору сложно определить доступных исполнителей, а руководитель организации не получает оперативную картину по мероприятиям, задачам и фактически подтвержденному вкладу. Дополнительно возрастает нагрузка на сотрудников, которые вручную проверяют заявки, считают часы, готовят справки и сертификаты.",
        "Актуальность темы выпускной квалификационной работы обусловлена необходимостью разработки специализированной web-системы, которая объединяет основные процессы волонтерской деятельности в одном интерфейсе и поддерживает как операционную работу координатора, так и мотивацию добровольцев. Для волонтера важны понятный личный кабинет, история участия, достижения и уведомления. Для координатора важны инструменты управления мероприятиями, задачами, заявками и отчетностью. Для администратора организации важны контроль доступа, аналитика и прозрачность действий пользователей.",
        "Объектом разработки является процесс управления волонтерской деятельностью в организации. Предметом разработки является web-ориентированная информационная система UnityAid, обеспечивающая учет волонтеров, организаций, мероприятий, задач, времени, достижений, сертификатов и аналитических показателей.",
        "Целью выпускной квалификационной работы является повышение эффективности управления волонтерской деятельностью за счет проектирования и разработки информационной системы UnityAid с элементами геймификации и аналитики.",
        "Для достижения цели необходимо решить следующие задачи:",
    ]
    for t in texts:
        paragraph(doc, t)
    number_list(doc, [
        "провести анализ предметной области и выявить проблемы существующей организации волонтерской деятельности;",
        "определить пользовательские роли, функциональные требования и ограничения системы;",
        "спроектировать архитектуру web-приложения, серверного API и базы данных;",
        "реализовать основные модули системы с разграничением доступа по ролям;",
        "проверить работоспособность приложения на ключевых пользовательских сценариях;",
        "оценить практическую значимость и направления дальнейшего развития системы.",
    ])
    paragraph(doc, "Методами выполнения работы являются системный анализ, моделирование предметной области, проектирование базы данных, проектирование REST API, разработка web-приложения, ручное и функциональное тестирование пользовательских сценариев.")
    paragraph(doc, "Практическая значимость работы состоит в создании программного продукта, который может использоваться как основа для внедрения в малых и средних волонтерских организациях, студенческих объединениях, благотворительных проектах и учреждениях, которым необходимы учет участников, прозрачная отчетность и повышение вовлеченности добровольцев.")


def chapter1(doc: Document):
    heading(doc, "1 ТЕОРЕТИЧЕСКИЙ АНАЛИЗ ПРЕДМЕТНОЙ ОБЛАСТИ И СУЩЕСТВУЮЩИХ РЕШЕНИЙ УПРАВЛЕНИЯ ВОЛОНТЕРСКОЙ ДЕЯТЕЛЬНОСТЬЮ", 1)
    heading(doc, "1.1 Анализ предметной области и информационных процессов управления добровольческой деятельностью", 2)
    for t in [
        "Волонтерская деятельность за последние годы приобрела особое значение как инструмент решения социальных, образовательных, культурных и организационных задач. В работе некоммерческих организаций, образовательных учреждений, муниципальных проектов и инициативных объединений волонтеры участвуют в мероприятиях, выполняют задачи, помогают координаторам и формируют устойчивое сообщество вокруг социально значимых целей.",
        "Рост числа добровольческих инициатив приводит к усложнению процессов управления. Организации должны учитывать сведения о волонтерах, их навыках и интересах, планировать мероприятия, распределять задачи, подтверждать участие, фиксировать затраченное время и формировать отчетность. При малом количестве участников эти процессы могут выполняться вручную, однако при увеличении масштаба ручная модель быстро становится источником ошибок и задержек.",
        "Основными участниками предметной области являются волонтеры, координаторы, администраторы организаций и руководители программ. Волонтеры подают заявки на мероприятия, выполняют поручения, фиксируют вклад и ожидают обратной связи. Координаторы организуют события, распределяют задачи, подтверждают участие и часы. Администраторы управляют пользователями, ролями и настройками организации. Руководители программ используют аналитические сведения для оценки эффективности деятельности.",
        "Для предметной области характерна необходимость централизованного учета. Данные о мероприятиях, заявках, задачах, часах и достижениях должны быть связаны между собой. Если сведения хранятся в разных таблицах, формах и переписках, организация теряет целостную картину: невозможно быстро определить активных участников, подтвердить вклад конкретного волонтера или подготовить объективный отчет за период.",
        "Отдельной проблемой является мотивация и удержание волонтеров. При отсутствии прозрачной системы признания участники не всегда видят результаты своей работы. Это снижает вовлеченность и затрудняет формирование постоянного сообщества. Элементы геймификации, такие как достижения, значки, уровни, рейтинги и персональная статистика, позволяют визуализировать вклад волонтера и поддерживать культуру признания заслуг.",
        "Развитие web-технологий и распространение мобильных устройств создают предпосылки для разработки единой платформы управления волонтерской деятельностью. Прогрессивное web-приложение, адаптивный интерфейс и клиент-серверная архитектура позволяют обеспечить доступ к системе как с рабочего компьютера координатора, так и со смартфона волонтера. Это особенно важно для молодежной аудитории и распределенных команд.",
    ]:
        paragraph(doc, t)
    table(doc, "Таблица 1 - Основные сущности предметной области", ["Сущность", "Назначение", "Пример данных"], [
        ["Волонтер", "Участник добровольческой деятельности", "ФИО, контакты, навыки, интересы, история участия"],
        ["Организация", "Пространство управления волонтерской программой", "Название, участники, роли, настройки, логотип"],
        ["Мероприятие", "Событие, требующее участия волонтеров", "Дата, место, формат, лимит, заявки, посещаемость"],
        ["Задача", "Поручение, назначаемое одному или нескольким участникам", "Описание, срок, приоритет, статус, исполнители"],
        ["Запись времени", "Фиксация затраченных волонтерских часов", "Дата, длительность, основание, статус подтверждения"],
        ["Достижение", "Элемент признания и геймификации", "Название, условие выдачи, баллы, пользователь"],
        ["Отчет", "Агрегированная управленческая информация", "Активность, часы, мероприятия, динамика, эффективность"],
    ], [3.2, 6.2, 6.6])
    heading(doc, "1.2 Сравнительный анализ существующих цифровых платформ для привлечения и координации волонтеров", 2)
    for t in [
        "Под сравнительным анализом цифровых платформ в рамках данной работы понимается не только сопоставление отдельных функций, но и оценка того, насколько существующие решения поддерживают полный цикл волонтерской деятельности. Для проектируемой системы важен путь от первичного привлечения участника до подтверждения его вклада: волонтер должен найти подходящее мероприятие, подать заявку, получить решение координатора, выполнить задачу или принять участие в событии, зафиксировать затраченное время и увидеть результат своей активности в профиле. Для организации, в свою очередь, значимы не отдельные формы регистрации, а управляемая среда, в которой сохраняются роли, история участия, отчеты, достижения и данные для последующей аналитики.",
        "Существующие решения закрывают этот процесс неравномерно. Одни платформы хорошо решают задачу публичного поиска добровольческих возможностей, но слабо поддерживают внутреннюю работу конкретной организации. Другие системы удобны для планирования задач или событий, однако не учитывают специфику добровольческой деятельности: подтверждение часов, мотивацию участников, выдачу сертификатов и накопление персональной истории волонтера. Поэтому анализ аналогов целесообразно проводить через вопрос о том, какую часть жизненного цикла волонтерской программы закрывает сервис и какие элементы приходится выносить во внешние инструменты.",
        "В качестве первой группы можно выделить универсальные CMS и конструкторы сайтов. С их помощью организация способна разместить описание проекта, опубликовать новости и собрать заявки через форму. Однако такая система остается в большей степени информационным сайтом: логика обработки заявок, разграничения ролей, учета часов, формирования достижений и построения аналитики должна создаваться отдельно. В результате организация получает набор доработок вокруг сайта, а не единую информационную систему управления добровольческой деятельностью.",
        "Корпоративные ERP- и ECM-платформы, напротив, обладают развитой процессной логикой, средствами документооборота и механизмами контроля исполнения. Их возможности полезны для крупных организаций, где требуется строгая регламентация и интеграция с другими корпоративными системами. Для небольших волонтерских объединений, образовательных учреждений и НКО такие решения часто оказываются избыточными: они требуют сложного внедрения, обучения пользователей и финансовых затрат, при этом повседневные сценарии волонтера остаются для них вторичными.",
        "Отдельную группу составляют сервисы управления проектами и регистрации на мероприятия. Они позволяют назначать задачи, фиксировать сроки, вести списки участников и контролировать статусы. Эти инструменты хорошо работают как вспомогательная среда координатора, но не формируют целостную добровольческую историю. В них отсутствует связь между заявкой на мероприятие, фактическим участием, подтвержденными часами, достижениями, сертификатами и отчетностью по активности. Именно эта связность данных является принципиальной для проектируемой системы.",
        "Наиболее близкими к предметной области являются специализированные добровольческие платформы. Для уточнения требований были рассмотрены реальные публичные решения: российская платформа Добро.рф, сервис Idealist, объединенный с VolunteerMatch, платформа Timecounts для организации волонтеров на событиях и система Golden для управления добровольцами. Эти сервисы показывают, какие подходы уже применяются в сфере добровольчества, но одновременно выявляют границы готовых решений относительно цели разработки UnityAid.",
    ]:
        paragraph(doc, t)
    paragraph(doc, "Платформа Добро.рф представляет собой крупную публичную экосистему добровольчества. Ее сильная сторона заключается в масштабе: пользователь получает доступ к каталогу инициатив, может искать проекты, участвовать в мероприятиях и вести личную траекторию волонтера. Такой подход особенно эффективен для массового привлечения участников и популяризации добровольческой деятельности. Вместе с тем для локальной организации платформа выступает прежде всего как внешняя экосистема, а не как полностью контролируемая информационная система, которую можно адаптировать под собственную структуру ролей, закрытые процессы, внутреннюю аналитику и правила геймификации.")
    add_screenshot(doc, IMG_DIR / "analogs" / "01_dobro.png", "Рисунок 1 - Главная страница платформы Добро.рф")
    paragraph(doc, "Idealist и связанный с ним VolunteerMatch ориентированы на поиск социальных проектов, вакансий и волонтерских возможностей. Пользовательский сценарий здесь строится вокруг подбора подходящей инициативы по тематике, месту и формату участия. Такая модель удобна для первичного контакта между добровольцем и организацией, однако после привлечения участника значительная часть операционной работы остается вне платформы: учет выполненных задач, подтверждение часов, формирование сертификатов и внутренняя мотивация требуют отдельных решений.")
    add_screenshot(doc, IMG_DIR / "analogs" / "02_volunteermatch.png", "Рисунок 2 - Публичная страница VolunteerMatch на платформе Idealist")
    paragraph(doc, "Timecounts ближе к задачам операционного управления. Сервис поддерживает рекрутинг, онбординг, расписания и работу с волонтерами на событиях. По сравнению с каталогами возможностей он сильнее ориентирован на деятельность координатора, что делает его полезным примером для анализа. При этом основная логика платформы сосредоточена вокруг событий и расписаний. Для UnityAid требуется более широкая модель, где мероприятие является только одной частью процесса наряду с задачами, подтверждением времени, достижениями, сертификатами и аналитикой по организации.")
    add_screenshot(doc, IMG_DIR / "analogs" / "04_timecounts.png", "Рисунок 3 - Страница платформы Timecounts")
    paragraph(doc, "Golden рассматривается как пример коммерческой SaaS-платформы управления волонтерами. В ней заметен акцент на автоматизацию коммуникаций, мобильный доступ и сопровождение добровольческих программ. Такой подход демонстрирует востребованность специализированных систем, но для дипломного проекта важна не покупка готового сервиса, а проектирование собственной архитектуры, модели данных и пользовательских сценариев. Кроме того, готовая внешняя платформа ограничивает контроль над структурой данных и возможностями дальнейшего развития системы под локальные процессы.")
    add_screenshot(doc, IMG_DIR / "analogs" / "05_golden.png", "Рисунок 4 - Страница платформы Golden")
    paragraph(doc, "Обобщение рассмотренных групп аналогов приведено в таблице 2. В таблице отражены не только преимущества каждого класса решений, но и ограничения, которые мешают использовать их как полноценную основу для UnityAid.")
    table(doc, "Таблица 2 - Сравнение групп аналогов", ["Группа решений", "Преимущества", "Ограничения для задачи UnityAid"], [
        ["Универсальные CMS", "Гибкость, расширения, быстрый запуск информационного сайта", "Нет встроенной логики заявок, часов, ролей, достижений и аналитики"],
        ["ERP/ECM-платформы", "Мощные инструменты документооборота и управления процессами", "Высокая стоимость, избыточность, сложность внедрения для небольших организаций"],
        ["Сервисы управления проектами", "Удобные задачи, исполнители, статусы и сроки", "Не учитывают мероприятия, волонтерские часы, сертификаты и мотивацию участников"],
        ["Сервисы мероприятий", "Регистрация участников и публичные страницы событий", "Недостаточно функций для долгосрочной истории волонтера и отчетности"],
        ["Добровольческие платформы", "Ориентация на добровольчество и поиск мероприятий", "Не всегда поддерживают автономное пространство организации, аналитику и геймификацию"],
    ], [4.0, 5.7, 6.3])
    paragraph(doc, "Более детальное сопоставление рассмотренных публичных платформ представлено в таблице 3. Оно показывает, что каждая из них решает важную часть задачи, но ни одна не закрывает одновременно локальное управление организацией, учет подтвержденного вклада, мотивационные элементы и аналитическую отчетность в рамках единой модели данных.")
    table(doc, "Таблица 3 - Сравнение рассмотренных цифровых платформ", ["Платформа", "Основная направленность", "Ограничения относительно цели разработки"], [
        ["Добро.рф", "Федеральная экосистема добровольчества: поиск проектов, участие в инициативах, личная траектория волонтера", "Платформа ориентирована на широкую публичную экосистему и не является локально развертываемой системой управления внутренними процессами отдельной организации"],
        ["Idealist / VolunteerMatch", "Поиск волонтерских возможностей, вакансий и организаций, связывание участников с социальными проектами", "Акцент сделан на публичном поиске и привлечении участников; функции внутреннего учета часов, задач, сертификатов и геймификации ограничены"],
        ["Timecounts", "Организация волонтеров для событий, рекрутинг, онбординг и расписания", "Система ближе к event-менеджменту и требует адаптации под многоорганизационную модель, аналитику и выдачу подтверждающих документов"],
        ["Golden", "Платформа управления волонтерами с автоматизацией, коммуникацией и мобильным приложением", "Коммерческое SaaS-решение, ориентированное на готовую платформу; в рамках дипломного проекта требуется собственная архитектура и контролируемая модель данных"],
    ], [3.4, 5.9, 6.7])
    for t in [
        "Проведенный анализ показывает, что UnityAid не должна повторять модель обычного каталога мероприятий или универсального трекера задач. Для достижения цели дипломного проекта требуется система, в которой организация получает собственное пространство управления, координатор работает с мероприятиями, заявками и задачами, а волонтер видит понятную историю своего участия. При этом каждое действие должно оставлять проверяемый след в данных: заявка переходит в участие, участие подтверждается часами, часы влияют на достижения, а накопленные сведения используются в отчетах.",
        "Таким образом, ключевым отличием UnityAid является связность пользовательских сценариев и данных. В системе заявка на мероприятие, посещаемость, запись времени, достижение и сертификат не существуют как разрозненные элементы. Они образуют единую цепочку, позволяющую отследить путь волонтера от регистрации до подтвержденного вклада и управленческой аналитики. Данный вывод используется далее при описании процесса управления волонтерской деятельностью и формировании функциональных требований к разрабатываемой системе.",
    ]:
        paragraph(doc, t)
    heading(doc, "1.3 Описание процесса организации мероприятий, обработки заявок и учета волонтерских часов", 2)
    for t in [
        "Процесс управления волонтерской деятельностью включает несколько взаимосвязанных этапов. На первом этапе организация формирует пространство работы: определяет координаторов, добавляет участников, настраивает роли и справочники. На втором этапе координатор планирует мероприятие или задачу, указывает сроки, описание, формат, ограничения и требования к участникам.",
        "После публикации мероприятия волонтеры подают заявки. Система должна сохранять статус каждой заявки, комментарии и дату подачи. Координатор рассматривает заявки, подтверждает участие, отклоняет неподходящие заявки или переводит часть участников в лист ожидания. Наличие статусов делает процесс прозрачным и позволяет участнику понимать свое текущее положение.",
        "Во время проведения мероприятия или выполнения задачи фиксируется фактическое участие. После завершения координатор отмечает посещаемость и подтверждает количество часов. Если волонтер самостоятельно создает запись времени, она должна пройти проверку. Только подтвержденные часы могут учитываться в профиле, аналитике, рейтингах и сертификатах.",
        "Следующий этап связан с мотивацией. На основании подтвержденных действий система может начислять достижения, баллы или уровни. Геймификация при этом должна быть встроена в реальные процессы, а не существовать отдельно от них. Достижение должно отражать конкретный вклад: участие, количество часов, выполнение задач, регулярность или инициативность.",
        "Завершающим этапом является аналитика и отчетность. Руководитель или администратор получает агрегированные показатели: количество активных волонтеров, проведенных мероприятий, выполненных задач, подтвержденных часов, выданных сертификатов и динамику активности по периодам. Эти сведения используются для планирования будущих мероприятий и оценки эффективности программ.",
    ]:
        paragraph(doc, t)
    add_picture(doc, IMG_DIR / "process.png", "Рисунок 5 - Организация процесса управления волонтерской деятельностью")
    table(doc, "Таблица 4 - Роли пользователей в процессе управления", ["Роль", "Функции в процессе", "Результат работы"], [
        ["Волонтер", "Подает заявки, выполняет задачи, фиксирует часы, получает достижения", "История участия, подтвержденные часы, сертификаты"],
        ["Координатор", "Создает мероприятия и задачи, рассматривает заявки, подтверждает участие", "Актуальные списки участников, корректные статусы, проверенные часы"],
        ["Администратор организации", "Управляет участниками, ролями, настройками и отчетами", "Настроенное пространство организации и контроль данных"],
        ["Руководитель программы", "Анализирует показатели и принимает управленческие решения", "Сводная аналитика по активности и эффективности"],
    ], [3.2, 7.2, 5.6])
    for t in [
        "Технический уровень процесса реализуется через централизованную базу данных, серверное API и web-интерфейс. Все операции выполняются через систему, поэтому данные остаются актуальными и связанными. Такой подход снижает количество ошибок, характерных для ручного переноса сведений между таблицами, чатами и документами.",
        "Для UnityAid важно поддерживать многоорганизационную модель. Один пользователь может состоять в нескольких организациях и иметь разные роли. Например, в одной организации он может быть координатором, а в другой - обычным волонтером. Поэтому права доступа должны проверяться с учетом контекста организации, а не только глобального признака пользователя.",
    ]:
        paragraph(doc, t)
    heading(doc, "1.4 Обоснование необходимости разработки специализированной информационной системы UnityAid", 2)
    for t in [
        "Необходимость разработки UnityAid обусловлена недостатками существующих решений и практическими потребностями небольших и средних волонтерских организаций. В реальной работе часто используются связки из таблиц, мессенджеров, онлайн-форм и ручных отчетов. Такая модель проста на старте, но плохо масштабируется и не обеспечивает достаточной надежности данных.",
        "Типовыми проблемами являются отсутствие централизованного учета волонтеров и их активности, сложность планирования мероприятий, невозможность быстро отследить историю участия, низкая прозрачность признания заслуг и нехватка аналитических инструментов. Координаторы тратят значительное время на ручную сверку списков, рассылку уведомлений и подготовку отчетности.",
        "UnityAid должна закрыть эти дефициты за счет единого web-приложения. В систему закладываются профили волонтеров, организации, мероприятия, задачи, учет времени, уведомления, база знаний, достижения, рейтинги, сертификаты и аналитика. Такой состав модулей отражает полный жизненный цикл волонтерской деятельности.",
        "Централизация данных обеспечивает хранение сведений в единой базе с учетом ролей и прав доступа. Специализированная бизнес-логика позволяет поддерживать сценарии, характерные именно для волонтерства: подачу заявки, рассмотрение координатором, подтверждение участия, начисление достижений, фиксацию времени и формирование отчетов.",
        "Многоорганизационный режим позволяет использовать одну систему для нескольких независимых организаций. Каждая организация получает собственных участников, мероприятия, задачи и настройки, а пользователь может работать в разных пространствах с разными ролями. Это важно для образовательных учреждений, сетей НКО и проектов с несколькими направлениями.",
        "Встроенная геймификация повышает вовлеченность участников. Значки, уровни и рейтинги не должны быть внешним декоративным элементом. Они должны опираться на подтвержденные действия и помогать волонтеру видеть накопленный вклад. Для координатора такая модель также полезна, поскольку позволяет выделять активных участников и поощрять регулярность.",
        "Развитая аналитика необходима для управленческих решений. Сводные панели позволяют отслеживать количество активных волонтеров, закрытых задач, проведенных мероприятий, отработанных часов и выданных сертификатов. Эти показатели помогают оценивать эффективность программ, планировать нагрузку и готовить отчеты для руководства и партнеров.",
        "Таким образом, разработка UnityAid обусловлена не только общей цифровизацией, но и конкретными дефицитами существующей практики управления волонтерской деятельностью. Специализированная web-ориентированная система позволяет объединить учет, координацию, мотивацию и аналитику в одном программном продукте.",
    ]:
        paragraph(doc, t)
    heading(doc, "1.5 Выводы по результатам анализа предметной области и существующих решений", 2)
    for t in [
        "В первой главе рассмотрены особенности современной волонтерской деятельности, основные участники процесса и проблемы, возникающие при ручной организации работы. Показано, что разрозненные таблицы, мессенджеры и отдельные формы не обеспечивают достаточной прозрачности, усложняют подтверждение вклада и увеличивают нагрузку на координаторов.",
        "Анализ предметной области позволил выделить базовые сущности, необходимые для эффективного управления: волонтеры, организации, мероприятия, задачи, записи времени, достижения и отчеты. Именно эти сущности должны стать ядром разрабатываемой информационной системы, поскольку они отражают полный жизненный цикл участия волонтера.",
        "Рассмотрение аналогов показало, что универсальные CMS, ERP- и ECM-платформы, сервисы управления проектами и платформы мероприятий закрывают только часть потребностей. Они либо требуют значительной доработки, либо являются избыточными, либо не поддерживают геймификацию, многоорганизационный режим и аналитику в нужном объеме.",
        "На основании выполненного анализа обоснована необходимость разработки специализированной системы UnityAid. Первая глава формирует теоретическую основу проекта и задает требования, которые далее конкретизируются в архитектуре, модели данных, серверной и клиентской реализации.",
    ]:
        paragraph(doc, t)


def chapter2(doc: Document):
    heading(doc, "2 ПРОЕКТИРОВАНИЕ И РАЗРАБОТКА ИНФОРМАЦИОННОЙ СИСТЕМЫ UNITYAID", 1)
    heading(doc, "2.1 Общая архитектура web-приложения", 2)
    for t in [
        "UnityAid спроектирована как web-приложение с разделением на клиентскую часть, серверное API и уровень хранения данных. Такой подход соответствует характеру задачи: пользователи работают через браузер, бизнес-логика сосредоточена на сервере, а данные хранятся в реляционной базе.",
        "Клиентская часть реализована на Vue.js 3 с использованием TypeScript, Vite и Vue Router. Она отвечает за маршрутизацию, отображение страниц, формы ввода, таблицы, фильтры, навигацию и взаимодействие пользователя с API. Серверная часть реализована на Go и предоставляет REST API для модулей авторизации, организаций, мероприятий, задач, учета времени, уведомлений, аналитики и администрирования.",
        "В качестве системы управления базами данных используется PostgreSQL. Выбор реляционной модели обусловлен большим количеством связей между пользователями, организациями, мероприятиями, задачами, заявками, часами и достижениями. Для локального запуска и проверки используется Docker Compose, что упрощает развертывание backend, frontend и базы данных в согласованной среде.",
        "Архитектурное разделение также упрощает дальнейшее масштабирование. При росте нагрузки клиентская часть может обслуживаться отдельным web-сервером или CDN, серверное API может масштабироваться горизонтально, а база данных может быть вынесена на отдельный управляемый сервер. Для дипломного проекта достаточно локального контейнерного запуска, но структура не препятствует развитию.",
        "REST API выбран как основной способ взаимодействия из-за простоты интеграции с web-интерфейсом и удобства отладки. Каждый предметный модуль имеет собственную группу маршрутов, а обмен данными выполняется в формате JSON. Такой подход понятен для frontend-разработки и хорошо соответствует CRUD-операциям над сущностями системы.",
        "Отдельный worker-процесс предусмотрен для операций, которые не должны блокировать основной запрос пользователя. К ним относятся отправка уведомлений, подготовка отчетов, начисление достижений после завершения мероприятия и иные фоновые задачи. Даже если часть таких функций реализуется постепенно, архитектура заранее учитывает необходимость асинхронной обработки.",
    ]:
        paragraph(doc, t)
    add_picture(doc, IMG_DIR / "architecture.png", "Рисунок 6 - Архитектура информационной системы UnityAid")
    heading(doc, "2.2 Функциональные требования и пользовательские роли", 2)
    for t in [
        "Ролевая модель является основой разграничения доступа в системе. Один пользователь может состоять в разных организациях, а права определяются не только глобальной системной ролью, но и членством внутри конкретной организации. Это позволяет использовать систему как для одной организации, так и как основу для дальнейшего multi-tenant развития.",
        "Неавторизованный пользователь имеет доступ только к входу, регистрации и восстановлению пароля. Волонтер получает доступ к личному кабинету, мероприятиям, задачам, часам, достижениям, уведомлениям и базе знаний. Координатор управляет операционными сущностями своей организации. Администратор организации управляет участниками и настройками. Суперадминистратор работает с системной административной панелью.",
        "Для каждой роли определяются не только разрешенные страницы, но и допустимые действия внутри страницы. Например, волонтер может видеть мероприятие и подать заявку, но не может подтвердить чужую заявку. Координатор может рассматривать заявки и отмечать посещаемость, но только в рамках организации, где он обладает соответствующей ролью.",
        "Состояния сущностей также являются частью требований. Мероприятие может быть черновиком, опубликованным, завершенным или отмененным. Заявка может ожидать рассмотрения, быть подтвержденной, отклоненной, отмененной или находиться в листе ожидания. Задача имеет собственный жизненный цикл от создания до завершения. Явное хранение статусов делает процесс управляемым и проверяемым.",
        "Для снижения ошибок пользовательский интерфейс должен подсказывать допустимые действия. Например, после завершения мероприятия координатору доступна отметка посещаемости и подтверждение часов, а действия по редактированию ключевых параметров должны быть ограничены. Это помогает сохранить целостность данных.",
    ]:
        paragraph(doc, t)
    table(doc, "Таблица 5 - Функциональные требования по ролям", ["Роль", "Основные функции", "Ограничения"], [
        ["Гость", "Регистрация, вход, восстановление пароля", "Нет доступа к рабочим разделам"],
        ["Волонтер", "Профиль, заявки, задачи, часы, достижения, уведомления", "Не может управлять чужими данными и настройками организации"],
        ["Координатор", "Мероприятия, заявки, задачи, посещаемость, подтверждение часов", "Работает только в доступных организациях"],
        ["Администратор организации", "Участники, роли, данные организации, отчеты", "Не управляет глобальными настройками приложения"],
        ["Суперадминистратор", "Системная админ-панель, организации, глобальные сущности", "Используется для технического администрирования"],
    ], [3.2, 7.0, 5.8])
    add_picture(doc, IMG_DIR / "modules.png", "Рисунок 7 - Функциональные модули системы UnityAid")
    heading(doc, "2.3 Проектирование базы данных", 2)
    for t in [
        "База данных проектировалась исходя из необходимости хранить как справочную, так и операционную информацию. Центральными сущностями являются users, organizations, organization_members, volunteer_profiles, events, tasks, time_entries, achievements, notifications, certificates и audit_log. Между ними существуют устойчивые связи, отражающие реальные процессы системы.",
        "Таблица users хранит учетные записи и базовые данные пользователя. Профиль волонтера вынесен в отдельную сущность, что позволяет отделить данные авторизации от предметных характеристик: навыков, интересов, контактной информации и истории участия. Организации связаны с пользователями через таблицу членства, где фиксируется роль участника внутри организации.",
        "Мероприятия и задачи привязаны к организациям. Заявки на мероприятия позволяют хранить статус участия, комментарии и решения координатора. Учет времени связывается с пользователем, организацией, мероприятием или задачей и проходит через подтверждение. Такой подход дает возможность формировать отчеты по разным срезам: по волонтеру, мероприятию, организации, периоду и статусу.",
        "При проектировании учитывалась необходимость восстановления истории. Поэтому для ряда сущностей используются служебные поля даты создания, обновления и мягкого удаления. Мягкое удаление позволяет скрыть запись из интерфейса, но сохранить возможность аудита и анализа, если запись уже участвовала в связанных процессах.",
        "Для справочников навыков, интересов, категорий новостей и категорий базы знаний используется отдельное хранение. Это позволяет администраторам постепенно настраивать систему под конкретную организацию, не меняя программный код. В перспективе такие справочники могут быть частично глобальными и частично локальными для организации.",
        "Индексы и внешние ключи должны обеспечивать быстрый поиск по наиболее частым сценариям: список мероприятий организации, задачи пользователя, записи времени за период, заявки по мероприятию, уведомления текущего пользователя. Это особенно важно для страниц со списками и фильтрами.",
    ]:
        paragraph(doc, t)
    table(doc, "Таблица 6 - Основные сущности базы данных", ["Сущность", "Назначение", "Основные связи"], [
        ["users", "Учетные записи и данные авторизации", "Связана с профилем, членством, задачами, часами, уведомлениями"],
        ["organizations", "Волонтерские организации и настройки", "Связана с участниками, мероприятиями, задачами, сертификатами"],
        ["events", "Мероприятия, смены, заявки и посещаемость", "Связана с организацией, заявками, отзывами, часами"],
        ["tasks", "Задачи и назначения исполнителям", "Связана с организацией, назначениями, комментариями, часами"],
        ["time_entries", "Записи волонтерских часов и статусы подтверждения", "Связана с пользователем, организацией, мероприятием или задачей"],
        ["achievements", "Правила достижений и выданные достижения", "Связана с пользователями и транзакциями баллов"],
        ["certificates", "Сертификаты и справки о волонтерской деятельности", "Связана с пользователем и организацией"],
        ["audit_log", "Журнал значимых действий пользователей", "Связан с пользователем и изменяемыми сущностями"],
    ], [3.7, 6.2, 6.1])
    heading(doc, "2.4 Проектирование серверной части", 2)
    for t in [
        "Серверная часть реализует бизнес-логику приложения и предоставляет REST API. Код организован по модулям: auth, users, organizations, events, tasks, timeentries, notifications, gamification, analytics, certificates, knowledge, admin и другим. Для каждого модуля выделяются доменные структуры, репозиторий, сервис и HTTP-обработчик.",
        "Разделение на слои повышает сопровождаемость проекта. Обработчик отвечает за прием HTTP-запроса и формирование ответа. Сервис содержит бизнес-правила: проверку прав, изменение статусов, начисление баллов, подтверждение часов. Репозиторий инкапсулирует SQL-запросы и работу с базой данных. Такой подход упрощает тестирование и снижает связность между транспортным уровнем и хранилищем.",
        "Авторизация основана на пользовательской сессии и проверке ролей. Для операций изменения данных сервер проверяет не только факт входа, но и принадлежность пользователя к организации, в рамках которой выполняется действие. Например, координатор может изменить заявку только для мероприятия своей организации, а администратор организации не получает автоматически системные права суперадминистратора.",
        "Важным элементом серверной части является аудит. Для операций создания, изменения и удаления записываются события, позволяющие восстановить историю действий. Это особенно важно для систем, где подтверждаются часы и формируются документы, поскольку организация должна понимать, кто и когда изменил статус записи.",
        "Обработка ошибок строится так, чтобы клиентская часть могла показать пользователю понятное сообщение. Ошибки валидации отделяются от ошибок доступа и внутренних ошибок сервера. Например, отсутствие обязательного поля должно возвращать один тип ответа, а попытка изменить данные чужой организации - другой.",
        "Для операций со списками используется постраничная выдача. Это предотвращает загрузку слишком большого объема данных в одном запросе и делает интерфейс стабильнее при росте числа пользователей, мероприятий или задач. Параметры поиска и фильтрации передаются через query-параметры API.",
        "Отдельное внимание уделяется сертификатам. Сервер формирует документ, сохраняет сведения о нем и присваивает проверочный код. Публичный endpoint проверки не раскрывает лишние персональные данные, но позволяет подтвердить подлинность документа и его связь с организацией.",
    ]:
        paragraph(doc, t)
    table(doc, "Таблица 7 - Основные группы API", ["Группа API", "Назначение"], [
        ["/auth", "Регистрация, вход, обновление сессии, выход, восстановление и смена пароля"],
        ["/organizations", "Создание организаций, управление участниками и ролями"],
        ["/events", "Мероприятия, заявки, посещаемость, смены и обратная связь"],
        ["/tasks", "Задачи, назначения, комментарии, вложения и история статусов"],
        ["/time-entries", "Создание, редактирование, подтверждение и отклонение часов"],
        ["/analytics", "Отчеты по активности, мероприятиям, задачам, геймификации и аудиту"],
        ["/certificates", "Генерация, скачивание и публичная проверка сертификатов"],
        ["/admin", "Универсальная системная административная панель"],
    ], [4.0, 12.0])
    heading(doc, "2.5 Проектирование клиентской части", 2)
    for t in [
        "Клиентская часть UnityAid построена как одностраничное приложение. Основная навигация размещена в боковом меню, а верхняя панель содержит элементы профиля, уведомлений и быстрого доступа. Такой интерфейс подходит для регулярной работы координаторов и администраторов, поскольку основные разделы доступны из единого рабочего пространства.",
        "Страницы приложения соответствуют предметным модулям: DashboardPage, EventsPage, EventDetailPage, TasksPage, TaskDetailPage, VolunteersPage, OrganizationsPage, TimeEntriesPage, AnalyticsPage, AchievementsPage, CertificatesPage, KnowledgeBasePage, NotificationsPage и AdminPanelPage. Для обмена с сервером используются отдельные API-слои внутри entities, что уменьшает дублирование запросов в компонентах.",
        "Формы ввода реализуются через переиспользуемые компоненты. Для списков используются постраничная навигация, фильтры и поиск. Для ролей с ограниченным доступом интерфейс скрывает недоступные действия, однако окончательная проверка прав остается на сервере. Это важно, поскольку клиентская проверка улучшает удобство, но не является механизмом безопасности.",
        "Особое внимание уделено адаптивности. Система должна быть доступна не только с рабочего компьютера координатора, но и с ноутбука или мобильного устройства волонтера. Поэтому интерфейс не опирается на сложные таблицы там, где пользователю достаточно карточек, списков и компактных форм.",
        "Для состояния авторизации используется централизованное хранилище. Это позволяет не передавать данные текущего пользователя вручную между страницами и единообразно реагировать на истечение сессии. При загрузке приложения выполняется проверка текущего пользователя, после чего маршрутизатор определяет доступность разделов.",
        "API-клиенты в папке сущностей отвечают за конкретные запросы к backend. Такой подход делает компоненты интерфейса проще: страница вызывает функцию, передает параметры и получает типизированный результат. При изменении endpoint достаточно обновить соответствующий API-модуль.",
        "Для страниц с большим количеством записей используются фильтры и пагинация. Например, список задач можно ограничить по статусу, организации или поисковой строке. Это важно для координаторов, которые ежедневно работают с оперативными списками и должны быстро находить нужную запись.",
        "Визуальный стиль интерфейса выбран рабочим и сдержанным. Для системы управления важнее читаемость, предсказуемая навигация и плотность полезной информации, чем декоративность. Основной экран должен помогать пользователю принять решение и выполнить действие за минимальное число шагов.",
    ]:
        paragraph(doc, t)
    heading(doc, "2.6 Реализация модулей аналитики, геймификации и сертификатов", 2)
    for t in [
        "Модуль аналитики предназначен для получения управленческой информации. Он агрегирует данные по волонтерам, мероприятиям, задачам, подтвержденным часам, начисленным достижениям и действиям пользователей. Наличие аналитики позволяет руководителю организации видеть не только отдельные записи, но и динамику активности.",
        "Геймификация реализуется через достижения, баллы и лидерборд. Система может начислять достижения за участие в мероприятиях, выполнение задач, накопление часов и регулярную активность. При этом геймификация рассматривается как вспомогательный механизм обратной связи, а не как замена содержательной мотивации волонтера.",
        "Модуль сертификатов формирует документы, подтверждающие участие и вклад пользователя. Сертификат связывается с организацией и пользователем, имеет код проверки и может быть скачан в PDF. Публичная проверка по коду позволяет внешнему получателю убедиться, что документ действительно был сформирован системой.",
        "Связка аналитики, геймификации и сертификатов повышает ценность системы. Волонтер получает видимый результат участия, координатор получает инструмент подтверждения вклада, а организация получает основание для отчетности и принятия управленческих решений.",
        "Аналитика строится на уже накопленных операционных данных, поэтому ее точность зависит от корректности процессов подтверждения. Неподтвержденные часы не должны попадать в итоговые показатели наравне с подтвержденными. Это правило делает отчеты надежнее и снижает риск завышения активности.",
        "Достижения проектируются как расширяемый справочник правил. В базовой версии достаточно достижений за количество часов, участие в мероприятиях и выполнение задач. В дальнейшем правила могут учитывать регулярность участия, направления волонтерства, качество обратной связи или работу в команде.",
        "Сертификаты и справки являются важной частью практической ценности системы. Для студента или волонтера документ подтверждает вклад, а для организации упрощает оформление результатов. Проверочный код снижает риск подделки документа и позволяет внешнему получателю выполнить самостоятельную проверку.",
    ]:
        paragraph(doc, t)
    table(doc, "Таблица 8 - Реализованные модули UnityAid", ["Модуль", "Результат реализации"], [
        ["Авторизация", "Регистрация, вход, refresh-сессии, выход, восстановление и смена пароля"],
        ["Организации", "Управление организациями, участниками, ролями, логотипом и данными"],
        ["Мероприятия", "Создание, редактирование, заявки, статусы, смены, завершение, посещаемость"],
        ["Задачи", "Создание задач, назначения, статусы, комментарии, вложения, история изменений"],
        ["Учет времени", "Создание записей, подтверждение и отклонение часов, связь с задачами и мероприятиями"],
        ["База знаний", "Статьи, категории, публикация, редактирование и поиск материалов"],
        ["Уведомления", "In-app уведомления о событиях и изменениях статусов"],
        ["Сертификаты", "Генерация, скачивание и публичная проверка документов"],
        ["Админ-панель", "Постраничное управление системными сущностями для суперадминистратора"],
    ], [4.0, 12.0])
    heading(doc, "2.6.1 Описание предметных модулей", 3)
    for t in [
        "Модуль авторизации обеспечивает начальную точку входа в систему. Он отвечает за регистрацию, вход, обновление сессии, выход, восстановление пароля, подтверждение email и смену пароля. Для остальных модулей он предоставляет сведения о текущем пользователе и его правах.",
        "Модуль пользователей и волонтерских профилей хранит сведения, необходимые для работы с участником. В профиле фиксируются контактные данные, навыки, интересы, история участия, подтвержденные часы и достижения. Разделение учетной записи и профиля позволяет не смешивать технические данные входа с предметной информацией.",
        "Модуль организаций реализует пространство, внутри которого работают участники. Организация объединяет волонтеров, координаторов и администраторов, а также является владельцем мероприятий, задач, новостей, сертификатов и аналитических показателей. Такая модель позволяет развивать систему в сторону нескольких независимых клиентов.",
        "Модуль мероприятий поддерживает создание событий, публикацию, подачу заявок, рассмотрение заявок, учет смен, отметку посещаемости и завершение. Он является одним из центральных модулей, потому что через мероприятия проходит значительная часть волонтерской активности.",
        "Модуль задач предназначен для распределения поручений между участниками. Задача имеет приоритет, срок, статус, исполнителей, комментарии, вложения и историю изменений. Это позволяет использовать систему не только для мероприятий, но и для постоянной операционной работы организации.",
        "Модуль учета времени связывает вклад волонтера с конкретным основанием: мероприятием, задачей или иной деятельностью внутри организации. Подтверждение часов координатором отделяет заявленные сведения от проверенных, что повышает достоверность отчетности.",
        "Модуль уведомлений информирует пользователя о значимых изменениях: статусе заявки, назначении задачи, подтверждении часов, появлении новых материалов или административных действиях. Наличие уведомлений снижает зависимость процессов от внешних мессенджеров.",
        "Модуль базы знаний позволяет организации хранить инструкции, правила участия, ответы на частые вопросы и обучающие материалы. Для новых волонтеров это снижает порог входа, а для координаторов уменьшает количество повторяющихся объяснений.",
        "Модуль аналитики формирует управленческие отчеты на основе операционных данных. Он помогает оценить активность волонтеров, эффективность мероприятий, выполнение задач, начисление достижений и действия пользователей. Аналитика делает систему полезной для руководителя, а не только для координатора.",
        "Модуль сертификатов завершает цикл подтверждения вклада. Документ формируется на основе проверенных данных, связывается с пользователем и организацией, получает код проверки и может быть передан внешнему получателю. Это повышает практическую ценность системы для образовательной и социальной среды.",
    ]:
        paragraph(doc, t)
    heading(doc, "2.6.2 Обеспечение безопасности и целостности данных", 3)
    for t in [
        "Безопасность UnityAid рассматривается на нескольких уровнях: защита учетных записей, разграничение доступа, валидация входных данных, контроль операций изменения и защита персональных сведений. Для системы, работающей с данными волонтеров, эти вопросы имеют практическое значение, поскольку в профилях могут храниться контакты, история участия и сведения о подтвержденных часах.",
        "Пароли пользователей не должны храниться в открытом виде. При регистрации и смене пароля сервер сохраняет только криптографический хэш. При входе пользователь передает пароль по защищенному соединению, сервер сравнивает его с хэшем и выдает сессию только при успешной проверке. Такой подход снижает последствия возможной компрометации базы данных.",
        "Разграничение доступа строится на проверке роли и принадлежности пользователя к организации. Это особенно важно для многоорганизационной модели. Пользователь может быть координатором в одной организации и обычным волонтером в другой, поэтому нельзя опираться только на глобальный признак роли. Каждая операция должна проверять контекст организации.",
        "Валидация данных выполняется как на клиенте, так и на сервере. Клиентская валидация помогает пользователю быстрее исправить ошибку в форме, но не считается достаточной для защиты. Сервер повторно проверяет обязательные поля, допустимые значения статусов, формат дат, ограничения по лимитам участников и права пользователя на изменение записи.",
        "Для сохранения целостности данных важно ограничивать переходы между статусами. Например, отмененное мероприятие не должно принимать новые заявки, завершенное мероприятие не должно произвольно менять дату проведения, а подтвержденные часы должны изменяться только пользователем с соответствующими правами. Явные правила переходов делают поведение системы предсказуемым.",
        "Журнал аудита фиксирует значимые операции: создание, изменение и удаление записей, изменение статусов, подтверждение часов, управление ролями. Наличие аудита повышает доверие к системе и помогает разбирать спорные ситуации. Если данные были изменены ошибочно, организация может установить, какое действие привело к проблеме.",
    ]:
        paragraph(doc, t)
    heading(doc, "2.6.3 Организация пользовательского интерфейса", 3)
    for t in [
        "Интерфейс UnityAid проектировался как рабочее пространство, а не как рекламная страница. Основная задача пользователя - быстро перейти к нужному разделу, увидеть актуальные данные и выполнить действие. Поэтому навигация строится вокруг постоянного бокового меню, а страницы используют повторяемые элементы: заголовок, фильтры, список, карточку деталей и форму.",
        "Для волонтера важны простые сценарии: увидеть доступные мероприятия, подать заявку, посмотреть назначенные задачи, добавить часы и получить уведомление о решении координатора. Эти действия не должны требовать знания внутренней структуры организации. Поэтому интерфейс волонтера скрывает административные детали и показывает только необходимые статусы.",
        "Для координатора интерфейс должен быть более плотным. Ему необходимо работать со списками заявок, участников, задач и часов. Поэтому используются таблицы, фильтры, постраничная навигация и быстрые действия. При этом формы должны оставаться понятными, чтобы координатор мог создавать мероприятие или задачу без обращения к технической документации.",
        "Для администратора организации важны настройки и контроль. На страницах организации отображаются участники, роли, основные сведения, логотип и управленческие действия. Такая структура позволяет разделить повседневную операционную работу координатора и более редкие административные задачи.",
        "Суперадминистраторская панель реализует универсальное управление сущностями. Она необходима не для ежедневной работы волонтерской организации, а для технической поддержки приложения, исправления данных и контроля системных справочников. Доступ к ней должен быть максимально ограничен.",
        "Адаптивность интерфейса учитывает разные устройства. Координатор чаще работает за компьютером, но волонтер может использовать телефон. Поэтому ключевые пользовательские действия должны корректно отображаться на узких экранах, а сложные административные таблицы должны сохранять читаемость на рабочем экране.",
    ]:
        paragraph(doc, t)
    table(doc, "Таблица 9 - Карта основных экранов клиентской части", ["Экран", "Назначение", "Основные действия"], [
        ["LoginPage", "Вход пользователя в систему", "Ввод email и пароля, переход к восстановлению доступа"],
        ["OnboardingPage", "Первичное знакомство и регистрационный сценарий", "Регистрация, выбор начального действия"],
        ["DashboardPage", "Сводная рабочая панель", "Просмотр ключевых показателей и быстрых переходов"],
        ["EventsPage", "Список мероприятий", "Фильтрация, просмотр, создание мероприятия координатором"],
        ["EventDetailPage", "Карточка мероприятия", "Подача заявки, рассмотрение заявок, отметка посещаемости"],
        ["TasksPage", "Список задач", "Фильтрация, создание, назначение и изменение статуса"],
        ["TaskDetailPage", "Карточка задачи", "Просмотр описания, комментарии, вложения, учет времени"],
        ["TimeEntriesPage", "Учет волонтерских часов", "Создание, редактирование, подтверждение и отклонение записей"],
        ["VolunteersPage", "Список волонтеров организации", "Поиск, фильтрация, просмотр профиля и активности"],
        ["OrganizationsPage", "Список организаций", "Просмотр доступных организаций и создание новой при наличии прав"],
        ["OrganizationDetailPage", "Карточка организации", "Управление участниками, ролями, логотипом и настройками"],
        ["AnalyticsPage", "Аналитика организации", "Просмотр отчетов по активности, мероприятиям, задачам и часам"],
        ["AchievementsPage", "Достижения и рейтинг", "Просмотр личных достижений и лидерборда"],
        ["CertificatesPage", "Сертификаты и справки", "Генерация, скачивание и просмотр документов"],
        ["CertificateVerifyPage", "Публичная проверка сертификата", "Проверка документа по коду без входа в систему"],
        ["KnowledgeBasePage", "База знаний", "Просмотр статей, поиск и переход к материалам"],
        ["KnowledgeFormPage", "Форма статьи базы знаний", "Создание и редактирование материала"],
        ["NewsListPage", "Новости организации", "Просмотр объявлений и публикаций"],
        ["NotificationsPage", "Уведомления пользователя", "Просмотр событий и отметка уведомлений прочитанными"],
        ["AdminPanelPage", "Системное администрирование", "Постраничное управление разрешенными сущностями"],
    ], [4.0, 5.5, 6.5])
    heading(doc, "2.6.4 Развертывание и сопровождение", 3)
    for t in [
        "Для локального запуска проекта используется Docker Compose. Он объединяет backend, frontend, базу данных и сопутствующие сервисы в единый сценарий запуска. Это снижает количество ручных действий и помогает воспроизводить окружение на другой машине.",
        "Миграции базы данных позволяют пошагово изменять структуру таблиц. Это важно для проекта, который развивается итерационно: при добавлении сертификатов, уведомлений или новых связей не требуется вручную изменять схему базы. Миграции фиксируют историю изменений и позволяют развернуть проект с нуля.",
        "Для production-развертывания предусмотрены отдельные Dockerfile и конфигурация прокси. Такой подход отделяет режим разработки от режима эксплуатации. В разработке важны быстрые перезапуски и удобство отладки, а в эксплуатации - стабильная сборка, предсказуемые переменные окружения и корректная маршрутизация запросов.",
        "Сопровождение системы предполагает резервное копирование базы данных, контроль логов, обновление зависимостей и проверку доступности сервиса. Даже если в рамках дипломной работы эти процессы описаны на уровне проектного решения, их учет показывает готовность архитектуры к реальному использованию.",
        "Важным направлением сопровождения является управление клиентскими экземплярами. UnityAid может развиваться как коробочная система для отдельной организации или как SaaS-платформа. В обоих случаях необходимы процедуры создания организации, настройки домена, загрузки логотипа, выдачи начальных ролей и проверки работоспособности.",
    ]:
        paragraph(doc, t)
    table(doc, "Таблица 10 - Эксплуатационные требования к системе", ["Требование", "Описание"], [
        ["Резервное копирование", "Регулярное сохранение базы данных и файловых вложений"],
        ["Логирование", "Фиксация ошибок приложения и значимых действий пользователей"],
        ["Мониторинг", "Контроль доступности API, базы данных и фоновых задач"],
        ["Обновление", "Применение миграций и поставка новых версий без потери данных"],
        ["Настройка клиента", "Создание организации, назначение ролей, загрузка логотипа и базовых справочников"],
    ], [5.0, 11.0])
    heading(doc, "2.7 Выводы по второму разделу", 2)
    for t in [
        "Во втором разделе описаны архитектура, функциональные требования, модель данных и реализация ключевых частей UnityAid. Система построена как web-приложение с клиентской частью на Vue.js, серверной частью на Go и базой данных PostgreSQL.",
        "Разработанная структура модулей обеспечивает поддержку основных процессов волонтерской организации: работу пользователей и ролей, мероприятия, задачи, учет времени, уведомления, базу знаний, достижения, аналитику, сертификаты и системное администрирование.",
    ]:
        paragraph(doc, t)


def chapter3(doc: Document):
    heading(doc, "3 ПРОВЕРКА РАБОТОСПОСОБНОСТИ И ОЦЕНКА РЕЗУЛЬТАТОВ", 1)
    heading(doc, "3.1 Организация испытаний", 2)
    for t in [
        "Проверка работоспособности UnityAid выполнялась по основным пользовательским сценариям. Цель испытаний состояла в подтверждении того, что реализованные модули корректно взаимодействуют друг с другом и обеспечивают полный цикл работы: от регистрации пользователя до формирования отчетных данных и сертификата.",
        "Backend проверяется запуском тестов, применением миграций на чистой базе и ручной проверкой API. Frontend проверяется сборкой проекта, переходами по основным страницам, заполнением форм и выполнением операций создания, редактирования и удаления сущностей. Отдельно проверяются права доступа, поскольку ошибки разграничения ролей могут привести к раскрытию или изменению чужих данных.",
        "Испытания целесообразно проводить на подготовленном демонстрационном наборе данных. В него входят несколько пользователей с разными ролями, одна или несколько организаций, мероприятия в разных статусах, задачи, записи времени, достижения и сертификаты. Такой набор позволяет быстро проверить не только пустые формы, но и реальные списки.",
        "Проверка backend начинается с миграций, потому что ошибки схемы базы данных проявляются во всех остальных модулях. После успешного применения миграций проверяются базовые endpoint авторизации, затем CRUD-операции и сценарии, где несколько сущностей взаимодействуют между собой.",
        "Проверка frontend включает не только сборку, но и ручной проход по интерфейсу. Важно убедиться, что пользователь видит только допустимые пункты меню, формы не отправляют некорректные данные, ошибки API отображаются понятным образом, а после успешного действия список или карточка обновляются.",
    ]:
        paragraph(doc, t)
    table(doc, "Таблица 11 - Основные проверки системы", ["Направление", "Проверяемые действия", "Ожидаемый результат"], [
        ["Авторизация", "Регистрация, вход, refresh, logout, смена пароля", "Пользователь получает доступ только после успешной проверки учетных данных"],
        ["Организации", "Создание организации, добавление участника, назначение роли", "Данные доступны в пределах организации и с учетом роли"],
        ["Мероприятия", "Создание мероприятия, подача и рассмотрение заявки, завершение", "Статусы заявок и мероприятия изменяются корректно"],
        ["Часы", "Создание записи, подтверждение, отклонение", "Подтвержденные часы учитываются в профиле и отчетах"],
        ["Геймификация", "Начисление достижений и баллов после активности", "Пользователь видит достижения и место в рейтинге"],
        ["Сертификаты", "Генерация, скачивание, публичная проверка", "Документ имеет проверочный код и доступен для проверки"],
        ["Аналитика", "Открытие отчетов и дашборда", "Отображаются агрегированные показатели по доступным данным"],
    ], [3.5, 6.2, 6.3])
    heading(doc, "3.2 Демонстрационный сценарий работы системы", 2)
    for t in [
        "Демонстрационный сценарий отражает типовую работу волонтерской организации. Сначала суперадминистратор открывает систему и создает организацию либо использует уже существующую. Затем администратор организации добавляет пользователей и назначает координатора. Координатор создает мероприятие, указывает формат, дату, место, лимит участников и публикует его для волонтеров.",
        "Волонтер входит в систему, открывает список мероприятий и подает заявку. Координатор рассматривает заявку и подтверждает участие. После проведения мероприятия координатор отмечает посещаемость и подтверждает часы. Система отражает часы в профиле пользователя, начисляет баллы и достижения, а затем позволяет сформировать сертификат.",
        "В завершающей части сценария администратор открывает аналитический раздел и проверяет показатели активности. Суперадминистратор может открыть журнал аудита и убедиться, что операции создания, изменения и удаления были зафиксированы. Такой сценарий подтверждает связность модулей и показывает практическую применимость системы.",
        "Сценарий также показывает, что система поддерживает разные уровни ответственности. Волонтер инициирует участие, но не может самостоятельно подтвердить часы. Координатор управляет мероприятием, но действует в пределах своей организации. Администратор видит более широкий контур и может управлять участниками. Суперадминистратор обслуживает систему в целом.",
        "При демонстрации важно проверять не только положительный путь, но и ограничения. Например, волонтер не должен видеть административную панель, координатор не должен изменять данные другой организации, а неподтвержденная запись времени не должна увеличивать итоговые часы в сертификате.",
        "Такая проверка позволяет выявить ошибки, которые не всегда обнаруживаются при изолированном тестировании отдельных endpoint. В реальной системе ценность создается именно связкой модулей, поэтому полный пользовательский сценарий является обязательной частью приемки.",
    ]:
        paragraph(doc, t)
    number_list(doc, [
        "суперадминистратор входит в систему и проверяет список организаций;",
        "администратор организации добавляет участника и назначает координатора;",
        "координатор создает мероприятие и публикует его;",
        "волонтер подает заявку на участие;",
        "координатор подтверждает заявку и отмечает посещаемость;",
        "волонтер или координатор добавляет запись времени;",
        "координатор подтверждает часы;",
        "система начисляет баллы и достижения;",
        "координатор формирует сертификат;",
        "проверяющий открывает публичную страницу проверки сертификата.",
    ])
    heading(doc, "3.3 Анализ результатов разработки", 2)
    for t in [
        "В результате выполнения работы получена информационная система, реализующая основные процессы управления волонтерской деятельностью. Разработка охватывает как операционные сценарии координатора, так и пользовательские сценарии волонтера. Наличие административной панели и аудита делает систему пригодной для использования не только как демонстрационного учебного проекта, но и как основы для практического внедрения.",
        "К сильным сторонам реализации относятся модульная структура backend, разделение API-слоя и интерфейса, использование PostgreSQL для связанных данных, наличие миграций, поддержка Docker Compose, а также широкий набор предметных модулей. Система не ограничивается регистрацией на мероприятия, а поддерживает задачи, часы, достижения, уведомления, базу знаний, сертификаты и аналитику.",
        "Ограничения текущей версии связаны с тем, что проект выполнялся в рамках выпускной квалификационной работы. Часть перспективных функций может быть расширена: интеграции с календарями и мессенджерами, расширенная настройка правил геймификации, полнотекстовый поиск, импорт больших списков волонтеров, развитая система шаблонов документов и мониторинг production-окружения.",
        "С точки зрения архитектуры результат можно считать расширяемым. Новые модули могут добавляться через отдельные группы маршрутов backend, новые таблицы миграций и новые страницы frontend. При этом уже существующая модель организаций и ролей позволяет сохранять общий принцип разграничения доступа.",
        "С точки зрения предметной области результат закрывает базовый управленческий цикл. Организация получает учет участников, мероприятий, задач и часов; волонтер получает личную историю и достижения; координатор получает инструменты ежедневной работы; руководитель получает аналитику и документы.",
        "С точки зрения учебной ценности проект демонстрирует применение нескольких компетенций: анализ предметной области, проектирование базы данных, разработка backend и frontend, работа с контейнерным окружением, проектирование интерфейса и подготовка документации.",
    ]:
        paragraph(doc, t)
    table(doc, "Таблица 12 - Соответствие задач работы полученным результатам", ["Задача", "Полученный результат"], [
        ["Анализ предметной области", "Выявлены роли, процессы, проблемы ручного учета и требования к системе"],
        ["Проектирование архитектуры", "Определена трехзвенная web-архитектура с REST API и PostgreSQL"],
        ["Проектирование базы данных", "Выделены сущности пользователей, организаций, мероприятий, задач, часов, достижений и сертификатов"],
        ["Разработка backend", "Реализованы модульные обработчики, сервисы и репозитории на Go"],
        ["Разработка frontend", "Созданы страницы и компоненты на Vue.js для основных пользовательских сценариев"],
        ["Проверка работоспособности", "Сформирована программа испытаний и проверены ключевые сценарии"],
    ], [6.0, 10.0])
    heading(doc, "3.4 Экономическая и практическая значимость", 2)
    for t in [
        "Практическая значимость UnityAid заключается в снижении трудозатрат координаторов на учет заявок, часов и отчетов. Если организация ведет процессы вручную, сотрудник регулярно тратит время на перенос данных между таблицами, сверку списков, подготовку справок и поиск истории участия. Единая система уменьшает число ручных операций и снижает риск ошибок.",
        "Экономический эффект для небольшой организации выражается прежде всего в экономии рабочего времени. Например, автоматическое формирование списков участников, подтверждение часов и подготовка сертификатов позволяют координатору уделять больше внимания содержательной работе с добровольцами. Для образовательных организаций это также повышает прозрачность учета внеучебной активности студентов.",
        "С точки зрения дальнейшего развития UnityAid может быть оформлена как коробочное решение для одной организации или как SaaS-платформа с изолированными пространствами организаций. Уже реализованные сущности организаций, членства и ролей создают основу для масштабирования продукта. Дополнительную ценность могут дать настройки бренда, публичные страницы организаций, конструктор процессов согласования и расширенные отчеты.",
        "Для оценки экономической целесообразности можно рассматривать сокращение времени на повторяющиеся операции. Если координатор еженедельно вручную сводит заявки, посещаемость и часы, автоматизация даже части этих действий дает заметную экономию. При увеличении числа мероприятий эффект растет, потому что система повторно использует уже введенные данные.",
        "Косвенный эффект выражается в повышении качества данных. Когда сведения вводятся один раз и используются в связанных модулях, уменьшается число расхождений между списками. Это особенно важно для отчетности, где ошибка в часах или участниках может привести к необходимости повторной проверки документов.",
        "Практическая ценность также связана с удержанием волонтеров. Прозрачная история участия, достижения, уведомления и сертификаты помогают участнику видеть результат своей работы. Для организации это означает более устойчивое сообщество и меньшие затраты на постоянное привлечение новых людей.",
    ]:
        paragraph(doc, t)
    table(doc, "Таблица 13 - Оценка сокращения ручных операций координатора", ["Операция", "Ручной подход", "Подход с UnityAid"], [
        ["Сбор заявок", "Форма, таблица и ручное уведомление участников", "Заявка создается в системе и получает статус"],
        ["Подтверждение участия", "Сверка списка и переписка с каждым участником", "Координатор меняет статус заявки в карточке мероприятия"],
        ["Учет часов", "Ручной подсчет по посещаемости и отдельным сообщениям", "Записи времени связаны с мероприятием или задачей"],
        ["Подготовка сертификата", "Ручное заполнение шаблона и проверка данных", "Документ формируется по подтвержденным данным"],
        ["Отчетность", "Сведение нескольких таблиц за период", "Показатели строятся на операционных данных системы"],
    ], [4.2, 5.8, 6.0])
    for t in [
        "Даже без точного коммерческого расчета видно, что основной экономический эффект связан с повторяемостью операций. Чем больше мероприятий проводит организация, тем чаще координатор выполняет однотипные действия. Автоматизация таких действий дает накопительный эффект и уменьшает зависимость качества учета от внимательности одного сотрудника.",
        "Дополнительная экономическая ценность связана с подготовкой документов. Если справки и сертификаты формируются вручную, сотрудник должен проверить ФИО, организацию, даты, количество часов и основание выдачи. В UnityAid эти данные уже связаны с пользователем, мероприятием и подтвержденными часами, поэтому риск ошибки ниже.",
        "Для студенческого волонтерства важен не только прямой экономический эффект, но и управленческая прозрачность. Университет или факультет может видеть активность по группам, направлениям и периодам, а студент получает подтверждение своего участия. Это повышает качество внеучебной работы и облегчает подготовку отчетности.",
        "Для благотворительных организаций система может быть полезна при подготовке отчетов перед партнерами. Наличие истории мероприятий, подтвержденных часов и сертификатов позволяет быстрее отвечать на вопросы о результатах программы и демонстрировать вклад добровольцев.",
        "Развитие UnityAid как продукта может идти поэтапно. На первом этапе система используется как локальный инструмент одной организации. На втором этапе добавляются настройки бренда, шаблоны документов и расширенная аналитика. На третьем этапе возможен переход к SaaS-модели, где каждая организация работает в собственном пространстве.",
        "В SaaS-модели важны биллинг, изоляция данных, управление тарифами, резервное копирование и мониторинг. Эти функции не являются обязательными для дипломного MVP, но текущая модель организаций, ролей и прав доступа создает основу для их последующего добавления.",
    ]:
        paragraph(doc, t)
    heading(doc, "3.5 Выводы по третьему разделу", 2)
    for t in [
        "В третьем разделе описана организация проверки работоспособности UnityAid, представлен демонстрационный сценарий и выполнена оценка результатов разработки. Проверка показывает, что система поддерживает полный рабочий цикл: создание организации, управление участниками, проведение мероприятия, подтверждение часов, начисление достижений, генерацию сертификата и просмотр аналитики.",
        "Полученный программный продукт имеет практическую значимость для автоматизации волонтерской деятельности и может развиваться в направлении коммерческого коробочного решения или multi-tenant SaaS-платформы.",
    ]:
        paragraph(doc, t)


def conclusion(doc: Document):
    heading(doc, "ЗАКЛЮЧЕНИЕ", 1)
    for t in [
        "В ходе выполнения выпускной квалификационной работы была спроектирована и разработана информационная система UnityAid, предназначенная для управления волонтерской деятельностью с элементами геймификации и аналитики. Работа включала анализ предметной области, постановку требований, проектирование архитектуры, разработку серверной и клиентской частей, проектирование базы данных и проверку основных сценариев.",
        "В первом разделе рассмотрены особенности волонтерской деятельности как объекта автоматизации. Выявлены проблемы ручного учета: разрозненность данных, сложность подтверждения часов, отсутствие прозрачной истории участия, высокая нагрузка на координаторов и недостаток аналитики. На основании анализа обоснована необходимость создания специализированной системы.",
        "Во втором разделе выполнено проектирование и описание реализации UnityAid. Система построена на клиент-серверной архитектуре: frontend реализован на Vue.js 3 и TypeScript, backend - на Go, база данных - PostgreSQL. В составе приложения реализованы модули авторизации, организаций, мероприятий, задач, учета времени, уведомлений, базы знаний, геймификации, аналитики, сертификатов и системного администрирования.",
        "В третьем разделе рассмотрена проверка работоспособности и оценка результатов. Демонстрационный сценарий подтвердил возможность прохождения полного цикла работы: от создания организации и мероприятия до подтверждения часов, начисления достижений и формирования сертификата. Полученные результаты соответствуют поставленной цели и задачам.",
        "Практическая значимость работы состоит в создании программной основы для автоматизации деятельности волонтерских организаций, студенческих объединений и социальных проектов. Дальнейшее развитие системы может быть связано с интеграциями календарей и мессенджеров, расширением аналитики, настройкой правил геймификации, импортом данных и развитием multi-tenant модели.",
    ]:
        paragraph(doc, t)


def references(doc: Document):
    heading(doc, "СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ", 1)
    refs = [
        "ГОСТ 7.32-2017. Система стандартов по информации, библиотечному и издательскому делу. Отчет о научно-исследовательской работе. Структура и правила оформления.",
        "ГОСТ Р 7.0.5-2008. Система стандартов по информации, библиотечному и издательскому делу. Библиографическая ссылка. Общие требования и правила составления.",
        "ГОСТ Р 2.105-2019. Единая система конструкторской документации. Общие требования к текстовым документам.",
        "СМК-О-СМГТУ-36-20. Выпускная квалификационная работа. Структура, содержание, общие правила выполнения и оформления.",
        "Выпускная квалификационная работа: от бакалавриата до аспирантуры: учебное пособие / О.С. Логунова, Л.Г. Егорова, М.Ю. Наркевич, Ю.В. Кочержинская. Магнитогорск: МГТУ им. Г.И. Носова, 2025.",
        "Федеральный закон от 11.08.1995 N 135-ФЗ «О благотворительной деятельности и добровольчестве (волонтерстве)».",
        "Федеральный закон от 27.07.2006 N 152-ФЗ «О персональных данных».",
        "Документация языка Go. URL: https://go.dev/doc/",
        "Документация Vue.js. URL: https://vuejs.org/guide/introduction.html",
        "Документация PostgreSQL. URL: https://www.postgresql.org/docs/",
        "Документация Docker. URL: https://docs.docker.com/",
        "Fielding R.T. Architectural Styles and the Design of Network-based Software Architectures. Doctoral dissertation. University of California, Irvine, 2000.",
        "Newman S. Building Microservices. Designing Fine-Grained Systems. O'Reilly Media, 2021.",
        "Fowler M. Patterns of Enterprise Application Architecture. Addison-Wesley, 2002.",
        "Martin R.C. Clean Architecture: A Craftsman's Guide to Software Structure and Design. Prentice Hall, 2017.",
        "OWASP Foundation. OWASP Top Ten Web Application Security Risks. URL: https://owasp.org/www-project-top-ten/",
        "Документация Vite. URL: https://vitejs.dev/guide/",
        "Документация TypeScript. URL: https://www.typescriptlang.org/docs/",
        "Документация OpenAPI Specification. URL: https://spec.openapis.org/oas/latest.html",
        "Документация JSON Web Tokens. URL: https://jwt.io/introduction",
        "Добро.рф. Официальный сайт платформы развития добровольчества. URL: https://dobro.ru/",
        "Idealist. VolunteerMatch is now part of Idealist. URL: https://www.idealist.org/volunteermatch",
        "Timecounts. Volunteer organizing platform. URL: https://www.timecounts.app/",
        "Golden. Volunteer management software. URL: https://goldenvolunteer.com/",
        "UnityAid. Техническое задание на разработку информационной системы управления волонтерами с элементами геймификации и аналитики. 2026.",
        "UnityAid. Описание реализованного функционала по модулям. 2026.",
        "UnityAid. Обзор API и ER-диаграмма. 2026.",
        "UnityAid. Программа и чеклист испытаний. 2026.",
        "UnityAid. Продуктовая стратегия развития системы. 2026.",
    ]
    for i, ref in enumerate(refs, 1):
        p = paragraph(doc, f"{i}. {ref}", first_indent=False)
        p.paragraph_format.left_indent = Cm(0)


def appendices(doc: Document):
    heading(doc, "ПРИЛОЖЕНИЕ А Фрагмент структуры базы данных", 1)
    paragraph(doc, "В приложении приведен укрупненный фрагмент логической структуры базы данных UnityAid, отражающий основные связи между сущностями системы.", first_indent=False)
    paragraph(doc, "USERS ||--o{ ORGANIZATION_MEMBERS : has", first_indent=False)
    paragraph(doc, "ORGANIZATIONS ||--o{ ORGANIZATION_MEMBERS : includes", first_indent=False)
    paragraph(doc, "ORGANIZATIONS ||--o{ EVENTS : owns", first_indent=False)
    paragraph(doc, "ORGANIZATIONS ||--o{ TASKS : owns", first_indent=False)
    paragraph(doc, "USERS ||--|| VOLUNTEER_PROFILES : profile", first_indent=False)
    paragraph(doc, "EVENTS ||--o{ EVENT_APPLICATIONS : receives", first_indent=False)
    paragraph(doc, "EVENTS ||--o{ EVENT_ATTENDANCE : tracks", first_indent=False)
    paragraph(doc, "TASKS ||--o{ TASK_ASSIGNMENTS : assigns", first_indent=False)
    paragraph(doc, "USERS ||--o{ TIME_ENTRIES : submits", first_indent=False)
    paragraph(doc, "USERS ||--o{ VOLUNTEER_ACHIEVEMENTS : earns", first_indent=False)
    paragraph(doc, "USERS ||--o{ NOTIFICATIONS : receives", first_indent=False)
    paragraph(doc, "USERS ||--o{ CERTIFICATES : receives", first_indent=False)
    paragraph(doc, "ORGANIZATIONS ||--o{ CERTIFICATES : issues", first_indent=False)
    heading(doc, "ПРИЛОЖЕНИЕ Б Программа испытаний", 1)
    paragraph(doc, "Программа испытаний включает проверку backend, frontend и демонстрационного сценария. Backend проверяется запуском тестов, применением миграций и ручной проверкой API. Frontend проверяется сборкой, навигацией, адаптивностью и выполнением основных операций в интерфейсе.", first_indent=False)
    number_list(doc, [
        "выполнить проверку миграций базы данных на чистом окружении;",
        "проверить регистрацию, вход, обновление сессии, выход и смену пароля;",
        "проверить CRUD организаций, мероприятий, задач, новостей и базы знаний;",
        "проверить подачу заявки на мероприятие и изменение ее статуса координатором;",
        "проверить создание, подтверждение и отклонение записей времени;",
        "проверить начисление достижений и отображение рейтинга;",
        "проверить генерацию, скачивание и публичную проверку сертификата;",
        "проверить отображение отчетов аналитики и журнала аудита;",
        "выполнить сборку frontend и проверить переходы по основному меню.",
    ])
    heading(doc, "ПРИЛОЖЕНИЕ В Фрагмент API системы", 1)
    paragraph(doc, "В приложении приведен укрупненный перечень основных endpoint, используемых клиентской частью UnityAid. Фактическая структура API может уточняться при развитии проекта, однако приведенная группировка отражает реализованную архитектуру модулей.", first_indent=False)
    table(doc, "Таблица В.1 - Основные endpoint REST API", ["Группа", "Метод и путь", "Назначение"], [
        ["Авторизация", "POST /api/v1/auth/register", "Регистрация нового пользователя"],
        ["Авторизация", "POST /api/v1/auth/login", "Вход пользователя в систему"],
        ["Авторизация", "GET /api/v1/auth/me", "Получение сведений о текущем пользователе"],
        ["Организации", "GET /api/v1/organizations", "Получение списка доступных организаций"],
        ["Организации", "POST /api/v1/organizations", "Создание организации пользователем с правами администратора"],
        ["Организации", "POST\n/api/v1/organizations/{id}/members", "Добавление участника в организацию"],
        ["Мероприятия", "GET /api/v1/events", "Получение списка мероприятий с фильтрами"],
        ["Мероприятия", "POST /api/v1/events", "Создание мероприятия координатором"],
        ["Мероприятия", "POST\n/api/v1/events/{id}/applications", "Подача заявки волонтером"],
        ["Мероприятия", "PATCH\n/api/v1/events/{id}/applications/\n{applicationId}", "Изменение статуса заявки координатором"],
        ["Задачи", "GET /api/v1/tasks", "Получение списка задач"],
        ["Задачи", "POST /api/v1/tasks", "Создание задачи"],
        ["Учет времени", "POST /api/v1/time-entries", "Создание записи волонтерских часов"],
        ["Учет времени", "PATCH\n/api/v1/time-entries/{id}/approve", "Подтверждение записи времени"],
        ["Аналитика", "GET /api/v1/analytics/overview", "Получение сводных показателей"],
        ["Сертификаты", "POST\n/api/v1/certificates/generate", "Генерация сертификата или справки"],
        ["Сертификаты", "GET\n/api/v1/certificates/verify/{code}", "Публичная проверка сертификата по коду"],
    ], [3.7, 5.5, 6.3], font_size=10)
    paragraph(doc, "Пример структуры запроса на создание мероприятия:", first_indent=False)
    for line in [
        "{",
        '  "organizationId": 1,',
        '  "title": "Помощь в проведении городского мероприятия",',
        '  "format": "offline",',
        '  "startsAt": "2026-06-10T10:00:00+05:00",',
        '  "endsAt": "2026-06-10T15:00:00+05:00",',
        '  "location": "Магнитогорск, площадка мероприятия",',
        '  "capacity": 30',
        "}",
    ]:
        paragraph(doc, line, first_indent=False)
    paragraph(doc, "Пример ответа API содержит идентификатор созданного мероприятия, текущий статус, сведения об организации и служебные даты создания и изменения. Клиентская часть использует эти данные для перехода на страницу мероприятия и обновления списка.", first_indent=False)
    heading(doc, "ПРИЛОЖЕНИЕ Г Описание алгоритмов обработки заявок и часов", 1)
    paragraph(doc, "Алгоритм обработки заявки на мероприятие:", first_indent=False)
    number_list(doc, [
        "Пользователь открывает опубликованное мероприятие и отправляет заявку.",
        "Сервер проверяет, что пользователь авторизован и состоит в допустимой организации.",
        "Сервер проверяет, что мероприятие не отменено и не завершено.",
        "Сервер проверяет, не существует ли уже активной заявки этого пользователя.",
        "Если лимит участников не превышен, заявке присваивается статус ожидания рассмотрения или подтверждения согласно настройкам мероприятия.",
        "Если лимит превышен, заявке может быть присвоен статус листа ожидания.",
        "Координатор рассматривает заявку и принимает решение.",
        "Пользователь получает уведомление об изменении статуса.",
    ])
    paragraph(doc, "Алгоритм подтверждения волонтерских часов:", first_indent=False)
    number_list(doc, [
        "Волонтер создает запись времени и указывает организацию, дату, длительность и основание.",
        "Система сохраняет запись со статусом ожидания подтверждения.",
        "Координатор открывает список записей времени своей организации.",
        "Координатор проверяет связь записи с мероприятием или задачей.",
        "При корректных данных координатор подтверждает запись.",
        "Подтвержденные часы включаются в профиль волонтера и аналитические отчеты.",
        "При некорректных данных координатор отклоняет запись с комментарием.",
        "Система фиксирует действие в журнале аудита и отправляет уведомление волонтеру.",
    ])
    paragraph(doc, "Алгоритм начисления достижения:", first_indent=False)
    number_list(doc, [
        "После подтверждения часов или завершения мероприятия система получает событие активности.",
        "Сервис геймификации загружает правила достижений, применимые к организации или системе в целом.",
        "Для каждого правила рассчитывается, выполнено ли условие пользователем.",
        "Если достижение еще не выдавалось и условие выполнено, создается запись о выдаче достижения.",
        "При необходимости создается транзакция баллов.",
        "Пользователь получает уведомление, а его профиль и рейтинг обновляются.",
    ])
    paragraph(doc, "Алгоритм генерации сертификата:", first_indent=False)
    number_list(doc, [
        "Координатор или администратор выбирает пользователя, организацию и период, за который требуется сформировать документ.",
        "Сервер проверяет права пользователя, запрашивающего генерацию сертификата.",
        "Система выбирает только подтвержденные записи времени за указанный период.",
        "Система рассчитывает суммарное количество часов и перечень оснований участия.",
        "Если подтвержденных данных недостаточно, сервер возвращает сообщение о невозможности сформировать документ.",
        "При успешной проверке создается запись сертификата с уникальным проверочным кодом.",
        "Сервер формирует PDF-документ и сохраняет сведения о нем в базе данных.",
        "Пользователь получает возможность скачать документ, а внешний проверяющий может открыть публичную страницу проверки по коду.",
    ])
    paragraph(doc, "Алгоритм формирования сводной аналитики:", first_indent=False)
    number_list(doc, [
        "Пользователь открывает аналитический раздел и выбирает организацию, период и тип отчета.",
        "Сервер проверяет, что пользователь имеет право просматривать аналитику выбранной организации.",
        "Система получает агрегированные данные по мероприятиям, задачам, заявкам, подтвержденным часам и достижениям.",
        "Неподтвержденные записи времени исключаются из итоговых показателей, но могут отображаться отдельно как ожидающие проверки.",
        "Показатели группируются по периоду, направлению, мероприятию или пользователю в зависимости от выбранного отчета.",
        "API возвращает структурированный результат, который frontend отображает в виде карточек, таблиц или графиков.",
        "При необходимости пользователь экспортирует результаты для дальнейшей подготовки управленческого отчета.",
    ])
    paragraph(doc, "Описанные алгоритмы позволяют отделить пользовательское действие от подтвержденного результата. Это особенно важно для учета часов и сертификатов, поскольку документ должен формироваться только на основании проверенных данных.", first_indent=False)
    paragraph(doc, "В проекте такой подход применяется последовательно: заявка еще не означает участие, созданная запись времени еще не означает подтвержденный вклад, а наличие активности еще не означает автоматическую выдачу документа. Между этими состояниями существуют проверки, выполняемые пользователями с соответствующими правами.", first_indent=False)
    paragraph(doc, "Разделение состояний повышает надежность отчетности. Если организация использует UnityAid для официального подтверждения волонтерской активности, то система должна хранить не только итоговое значение часов, но и путь его получения: мероприятие или задачу, пользователя, дату, статус и координатора, подтвердившего запись.", first_indent=False)
    paragraph(doc, "Таким образом, алгоритмы обработки заявок, часов, достижений, сертификатов и аналитики образуют единый контур управления данными. Каждый следующий этап использует результат предыдущего только после проверки, что снижает риск ошибочных начислений и повышает доверие к итоговым отчетам.", first_indent=False)


    add_screenshots_appendix(doc)
    add_code_appendix(doc)


def main():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    make_images()
    doc = Document(EXAMPLE_TEMPLATE)
    clear_document_body(doc)
    configure_section(doc.sections[0], page_numbers=True, start=4)
    set_styles(doc)
    contents_page(doc, add_break=False)
    introduction(doc)
    chapter1(doc)
    chapter2(doc)
    chapter3(doc)
    conclusion(doc)
    references(doc)
    appendices(doc)
    normalize_black_text(doc)
    doc.core_properties.author = STUDENT
    doc.core_properties.title = TITLE
    doc.core_properties.subject = "Выпускная квалификационная работа"
    doc.save(DOCX_PATH)
    print(DOCX_PATH)


if __name__ == "__main__":
    main()
