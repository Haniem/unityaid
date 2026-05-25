from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Mm, Pt


BASE_DIR = Path(__file__).resolve().parent
MD_PATH = BASE_DIR / "Краткое_объяснение_диплома.md"
DOCX_PATH = BASE_DIR / "Краткое_объяснение_диплома.docx"


def set_run(run, size=12, bold=False):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.bold = bold


def add_paragraph(doc, text="", style=None, bold=False, size=12, align=None):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15
    if align is not None:
        p.alignment = align
    if text:
        run = p.add_run(text)
        set_run(run, size=size, bold=bold)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    set_run(run, size=12)


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    set_run(run, size=12)


def main():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.left_margin = Mm(25)
    section.right_margin = Mm(15)
    section.top_margin = Mm(20)
    section.bottom_margin = Mm(20)

    styles = doc.styles
    for name in ["Normal", "List Bullet", "List Number"]:
        styles[name].font.name = "Times New Roman"
        styles[name]._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        styles[name].font.size = Pt(12)

    current_numbered = False
    for raw in MD_PATH.read_text(encoding="utf-8").splitlines():
        line = raw.strip()
        if not line:
            current_numbered = False
            continue
        if line.startswith("# "):
            add_paragraph(doc, line[2:], bold=True, size=16, align=WD_ALIGN_PARAGRAPH.CENTER)
        elif line.startswith("## "):
            add_paragraph(doc, line[3:], bold=True, size=13)
        elif line.startswith("- "):
            add_bullet(doc, line[2:])
        elif len(line) > 3 and line[0].isdigit() and line[1:3] == ". ":
            current_numbered = True
            add_number(doc, line[3:])
        else:
            text = line.replace("**", "")
            add_paragraph(doc, text)

    doc.core_properties.title = "Краткое объяснение темы диплома"
    doc.core_properties.author = "Зозин Павел Алексеевич"
    doc.save(DOCX_PATH)
    print(DOCX_PATH)


if __name__ == "__main__":
    main()
